"""Request 6: colleague-facing freeway-control review, through 2 October 2026.

Fourteen slides: two boilerplate and twelve substantive slides.
Native grouped diagrams; typeset SVG mathematics with embedded PNG fallbacks.
The layout proof uses DejaVu Sans, not an Office/Calibri rendering.
"""
from pathlib import Path
from io import BytesIO
import json
import math
import os
import posixpath
import re
import zipfile

os.environ.setdefault("MPLCONFIGDIR", "/tmp/freeway_slides_matplotlib")
os.environ.setdefault("XDG_CACHE_HOME", "/tmp/freeway_slides_font_cache")
Path(os.environ["XDG_CACHE_HOME"]).mkdir(parents=True, exist_ok=True)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
OUT = HERE / "weekly_freeway_control_review_20261002.pptx"
ASSETS = HERE / "weekly_freeway_control_review_20261002_assets"
NOTES = HERE / "weekly_freeway_control_review_20261002_notes.md"
REPORT = HERE / "weekly_freeway_control_review_20261002_validation.json"
PREVIEW = HERE / "weekly_freeway_control_review_20261002_preview.png"
ASSETS.mkdir(exist_ok=True)
FONT = "Calibri"
INK, TEAL, BLUE, AMBER = "243341", "267D86", "416F9B", "A87532"
GRAY, LIGHT, ROAD, WHITE = "7D8791", "F1F4F6", "E4E8EB", "FFFFFF"
EMU = 914400
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
VSL = "misc/differentiable_freeway_bottleneck_task_review"
RAMP = "misc/selective_discreteness_ramp_metering_review"
svg_sequence = 0
records = []
math_lines = {}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = "Freeway control: what traffic information changes useful decisions?"
prs.core_properties.subject = "Method-first VSL and ramp-metering review, through 2 October 2026"
prs.core_properties.author = "Research review"
prs.core_properties.keywords = "VSL, ramp metering, hidden information, useful gradients"
plt.rcParams.update({"svg.fonttype": "path", "mathtext.fontset": "cm",
                     "font.family": "DejaVu Sans"})


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, x, y, w, h, value, *, role="body", bold=False,
         color=INK, center=False, middle=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = role + ":" + value.replace("\n", " / ")
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(.035)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
    for i, line in enumerate(value.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.font.name = FONT
        p.font.size = Pt(32 if role == "title" else 18)
        p.font.bold = bold or role == "title"
        p.font.color.rgb = rgb(color)
        p.space_before = p.space_after = Pt(0)
    return shape


def rect(slide, x, y, w, h, fill=LIGHT, line=None, kind=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.3)
    else:
        s.line.fill.background()
    return s


def arrow(slide, x, y, w=.5, h=.22, color=GRAY, kind=MSO_SHAPE.RIGHT_ARROW):
    return rect(slide, x, y, w, h, color, kind=kind)


def line(slide, x1, y1, x2, y2, color=GRAY, width=1.7, end=False):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if end:
        e = OxmlElement("a:tailEnd")
        e.set("type", "triangle")
        s._element.spPr.get_or_add_ln().append(e)
    return s


def group(slide, start, name):
    s = slide.shapes.add_group_shape(list(slide.shapes)[start:])
    s.name = "diagram:" + name
    return s


def slide(title, explanation, sources, *, boilerplate=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    text(s, .65, .4, 12.02, .72, title, role="title")
    source_text = "\n".join(sources)
    s.notes_slide.notes_text_frame.text = explanation + "\n\nSources:\n" + source_text
    records.append({"slide": len(prs.slides), "title": title,
                    "boilerplate": boilerplate, "notes": explanation,
                    "sources": sources})
    return s


def body(slide, values, y=5.1, step=.48, x=.82, w=11.75):
    for i, value in enumerate(values):
        text(slide, x, y + step*i, w, .40, value)


def vector_picture(slide, name, svg, png, x, y, w, h):
    """Verified Part/relationship API; SVG blip plus normal PNG image part."""
    global svg_sequence
    with Image.open(png) as image:
        iw, ih = image.size
    factor = min(w/iw, h/ih)
    ww, hh = iw*factor, ih*factor
    pic = slide.shapes.add_picture(str(png), Inches(x+(w-ww)/2),
                                  Inches(y+(h-hh)/2), Inches(ww), Inches(hh))
    pic.name = name
    svg_sequence += 1
    part = Part(PackURI(f"/ppt/media/weekly_vector_{svg_sequence}.svg"),
                "image/svg+xml", slide.part.package, Path(svg).read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.find(qn("p:blipFill")).find(qn("a:blip"))
    extlist = blip.find(qn("a:extLst"))
    if extlist is None:
        extlist = OxmlElement("a:extLst")
        blip.append(extlist)
    ext = OxmlElement("a:ext")
    ext.set("uri", SVG_EXT)
    node = etree.Element(f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS})
    node.set(qn("r:embed"), rid)
    ext.append(node)
    extlist.append(ext)
    return pic


def formula(slide, name, lines, x, y, w, h, fontsize=25, line_spacing=.56):
    fig = plt.figure(figsize=(11.6, line_spacing*len(lines)+.18), dpi=240)
    fig.patch.set_alpha(0)
    for i, value in enumerate(lines):
        fig.text(.5, 1-(i+.6)/len(lines), "$"+value+"$",
                 ha="center", va="center", fontsize=fontsize, color="#"+INK)
    svg, png = ASSETS / (name+".svg"), ASSETS / (name+".png")
    fig.savefig(svg, transparent=True, bbox_inches="tight", pad_inches=.045)
    fig.savefig(png, transparent=True, bbox_inches="tight", pad_inches=.045)
    plt.close(fig)
    shape = vector_picture(slide, "math:"+name, svg, png, x, y, w, h)
    math_lines[shape.name] = len(lines)


def chain(slide, labels, *, y=1.62):
    start = len(slide.shapes)
    span, gap = 11.8, .35
    width = (span-gap*(len(labels)-1))/len(labels)
    for i, label in enumerate(labels):
        x = .78 + i*(width+gap)
        text(slide, x, y, width, .77, label, role="plot",
             bold=True, color=TEAL, center=True, middle=True)
        if i < len(labels)-1:
            arrow(slide, x+width+.025, y+.30, gap-.05, color=TEAL)
    group(slide, start, "causal sensitivity chain")


# Exactly two boilerplate slides.
s = slide("Freeway control and useful traffic gradients",
          "Weekly colleague presentation requested in feedback item 6. Scope starts "
          "with VSL task definition and ends with the current method-first ramp review. "
          "The VSL result is a closed conditional analytical certificate; the ramp "
          "information/action test remains open. This is presentation synthesis, not a new review or experiment.",
          [VSL+"/reporting/s13/s13_99_m2_certificate_sprint_closure.md",
           RAMP+"/reporting/rm9_current_handoff.md"], boilerplate=True)
text(s, .86, 2.05, 11.6, .75, "What information changes a useful control decision?",
     bold=True, color=TEAL)
body(s, ["Variable speed limits at a lane drop",
         "Ramp metering at a local freeway merge",
         "Methodology and findings • 2 October 2026"], y=3.05, step=.72)

s = slide("Contents",
          "The order mirrors the requested syllabus. Method-first principles and "
          "research questions frame both tasks. VSL includes its causal mechanism, "
          "exact conditional information/action result, gradient and study boundary. "
          "Ramp metering receives the same treatment where evidence exists, with open claims clearly identified.",
          ["temp_content/feedback/feedback.md"], boilerplate=True)
body(s, ["1. Review principle and core questions",
         "2. VSL task, hidden information and decision mechanism",
         "3. VSL gradient and research-readiness boundary",
         "4. Ramp-metering task, mechanism and information test",
         "5. Ramp gradient and the current evidence boundary"], y=1.8, step=.80)

s = slide("Start from the control question",
          "The historical implementation-anchored approach sought a compatible "
          "freeway task. The independent reviews instead start with a task and a strong "
          "aggregate comparator, then ask whether unavailable information changes "
          "useful actions. Current project code is not a model-family admission constraint. "
          "This does not claim existing simulator implementation was invalid.",
          [VSL+"/01_review_recap_through_s06.md",
           RAMP+"/reporting/rm7_20_method_first_task_formulation_and_gate.md"])
start = len(s.shapes)
text(s, .83, 1.45, 11.65, .4, "Earlier project anchoring", role="plot", bold=True, color=GRAY)
text(s, .83, 4.03, 11.65, .4, "Method-first review", role="plot", bold=True, color=TEAL)
for row, labs, col in [(2.24, ["Existing model", "Compatible task", "Evaluate"], GRAY),
                       (4.82, ["Task +\nstrong baseline", "Information value?", "Minimum state\n+ gradients"], TEAL)]:
    for i, label in enumerate(labs):
        x = .86 + 4.04*i
        rect(s, x, row, 3.36, .82, LIGHT, col)
        text(s, x+.05, row+.10, 3.26, .62, label, role="plot",
             center=True, middle=True, bold=True, color=col)
        if i<2:
            arrow(s, x+3.50, row+.29, .40, color=col)
group(s, start, "two research starting points")
body(s, ["Let decision evidence justify traffic detail."], y=6.38)

s = slide("Three questions guide both control tasks",
          "These are the persistent core review questions. Information loss means "
          "a difference unavailable to the complete declared aggregate observation "
          "history, not merely different lane snapshots. All comparisons use the same "
          "action family, known forcing and complete population objective. Useful "
          "gradient evidence differs from differentiability or lower internal loss.",
          [VSL+"/reporting/s13/s13_35_study_sufficiency_and_extension_audit.md",
           RAMP+"/synthesis/rm8_predecessor_and_information_map.md"])
body(s, ["1. What traffic information can be hidden from aggregate history?",
         "2. Can that information change a useful control action?",
         "3. Can retaining it provide relevant gradient information?"],
     y=1.80, step=1.0)
body(s, ["The required traffic resolution is an outcome of the review.",
         "Match actions, known forcing and complete travel/waiting accounting."], y=5.54, step=.6)

s = slide("VSL task: a permanent three-to-two lane drop",
          "One shared lane-uniform infrastructure advice speed is varied over a "
          "bounded interval. The selected analytical branch converts advice to traversal "
          "time; this is not an assertion of universal field compliance. The permanent "
          "drop removes reopening/recovery bookkeeping. Keeping one actuator avoids "
          "confounding information value with extra lane-specific control authority. "
          "The final proof uses ending-versus-continuing groups under continuing-group symmetry.",
          [VSL+"/synthesis/s13/s13_m2_model_contract.tex",
           VSL+"/extraction/s13_m2_source_to_equation_audit.md"])
start = len(s.shapes)
for y in [2.02,2.87]:
    rect(s, 1.0,y,11.2,.59,ROAD)
    arrow(s, 11.50,y+.18,.45)
rect(s,1.0,3.72,7.18,.59,ROAD)
line(s,8.18,4.015,9.45,3.165,end=True,width=2)
rect(s,3.15,1.80,3.32,2.75,None,TEAL)
text(s,2.81,1.30,4.08,.38,"One common speed advice",role="plot",center=True,bold=True,color=TEAL)
text(s,3.17,4.56,3.3,.38,"Preparation zone",role="plot",center=True)
text(s,8.55,4.12,3.42,.38,"Mandatory merge",role="plot",center=True)
group(s,start,"lane-drop geometry and common control")
body(s,["Permanent narrowing gives the smallest clear mandatory-merge task.",
        "One action for all lanes isolates information value from extra control.",
        "Outcome: complete cohort travel and waiting time through clearance."],y=5.26,step=.53)

s = slide("Preparation time can change receiving work",
          "In the accepted fixed smooth branch, slower advice increases L/v and "
          "permits more mandatory ending-lane mass to transfer before the choke. "
          "The constant-rate exponential is an explicit analytical adaptation. "
          "Remaining mandatory-merge mass consumes alpha>1 receiving-work units "
          "per vehicle-equivalent under proportional homogeneous batch service. "
          "Receiving work is not additional vehicle mass. All action-induced traversal "
          "time and complete clearance are charged. If service is independent of action, "
          "the earlier fixed-service obstruction restores a common fastest action.",
          [VSL+"/synthesis/s13/s13_m2_model_contract.tex",
           VSL+"/reporting/s13/s13_20_m2_certificate_findings_and_gate.md"])
start=len(s.shapes)
for y in [1.68,2.22,2.76]:
    rect(s,1.0,y,10.8,.35,ROAD)
for x in [2.2,4.1,5.95]:
    arrow(s,x,2.51,.3,.48,color=TEAL,kind=MSO_SHAPE.UP_ARROW)
rect(s,1.3,2.82,1.3,.22,AMBER)
rect(s,4.7,2.82,.8,.22,AMBER)
rect(s,7.9,2.82,.35,.22,AMBER)
line(s,8.8,2.935,9.5,2.395,end=True,width=2)
text(s,2.3,3.30,4.28,.38,"More time → more early merging",role="plot",center=True)
text(s,8.14,3.30,3.63,.38,"Less merge work at choke",role="plot",center=True)
group(s,start,"progressive lateral transfer before choke")
formula(s,"vsl_transfer",[r"\tau=\frac{L}{v}",r"e(\tau)=pM\,e^{-\lambda\tau}"],.9,3.93,11.5,1.13)
body(s,["L: zone length; v: common advice speed; τ: time in the zone.",
        "p: ending-lane fraction; M: cohort mass; λ: transfer rate; e: remaining mass.",
        "Slowing costs travel time; early merging can save receiving time.",
        "Fixed action-independent service instead favors fastest advice."],y=5.28,step=.45)

s = slide("Hidden allocation can change the best advice",
          "This is the accepted conditional VSL certificate, not a field finding. "
          "The two ending-lane fractions inherit exact predecision density, flow, speed "
          "and count history equality from the homogeneous conservative upstream construction. "
          "The spatial decision boundary is first entry to the common preparation zone. "
          "For the open parameter regime, unique actions differ over the entire speed interval: "
          "low p selects fastest, high p an interior slower action. Aggregate-history-only "
          "initial action has positive worst-state regret; later feedback may see divergence. "
          "The plot reproduces the certificate's normalized illustration: lambda=1, "
          "tau range [0.2,2], K=4, p_low=0.2, p_high=0.5. Curves are excess analytical "
          "mean time relative to each state's optimum, not simulation or calibration.",
          [VSL+"/synthesis/s13/s13_m2_certificate_analysis.tex",
           VSL+"/reporting/s13/s13_99_m2_certificate_sprint_closure.md"])
fig,ax=plt.subplots(figsize=(11.3,3.1),dpi=200)
tau=np.linspace(.2,2,300)
for p,col,label in [(.2,TEAL,"Less ending-lane mass"),(.5,AMBER,"More ending-lane mass")]:
    opt=max(.2,min(2,math.log(4*p)))
    regret=tau+4*p*np.exp(-tau)-(opt+4*p*math.exp(-opt))
    ax.plot(tau,regret,color="#"+col,lw=2.7,label=label)
    ax.scatter([opt],[0],color="#"+col,s=75,zorder=4)
    ax.axvline(opt,color="#"+col,lw=1.1,linestyle="--",alpha=.8)
ax.set_xlim(.16,2.03); ax.set_ylim(-.045,1.32)
ax.set_xlabel("Time in preparation zone (normalized)",fontsize=18)
ax.set_ylabel("Extra mean time",fontsize=18)
ax.tick_params(labelsize=16)
ax.spines[["top","right"]].set_visible(False)
ax.legend(frameon=False,fontsize=17,loc="upper left")
fig.tight_layout()
svg,png=ASSETS/"vsl_action_contrast.svg",ASSETS/"vsl_action_contrast.png"
fig.savefig(svg,bbox_inches="tight");fig.savefig(png,bbox_inches="tight");plt.close(fig)
vector_picture(s,"diagram:conditional analytical action contrast",svg,png,.88,1.36,11.68,3.44)
body(s,["Different lane fractions; identical declared predecision aggregate history.",
        "Less ending-lane mass: fastest advice. More: a slower interior optimum.",
        "One aggregate-history-only initial action cannot be optimal for both.",
        "Conditional analytical illustration; normalized values, not traffic data."],y=5.02,step=.50)

s = slide("VSL: follow the causal gradient path",
          "Fixed-active-regime chain: advice speed -> traversal time -> remaining "
          "mandatory-transfer mass -> receiving workload -> complete mean time. "
          "The displayed cost and derivative are exact for the proportional fluid batch. "
          "Hidden p changes e and can reverse the locally improving speed direction "
          "near fastest advice. Structural lane topology is not differentiated. "
          "Lane transfer structure is grounded in Nagalur Subraveti/Knoop/van Arem; "
          "gap-weighted receiving and proportional sharing in Carey/Balijepalli/Watling; "
          "common VSL lane-drop setting in Zhang et al. The exponential active branch, "
          "batch handoff and information/action/gradient certificate are review adaptations "
          "and deductions, not jointly established by one literature work. No AD/FD or "
          "independent reference evaluation occurred. J is total cohort system time "
          "divided by fixed cohort mass M, not last-departure time. All mass travels "
          "for tau and is then discharged as a homogeneous proportional batch. "
          "W=(M-e)+alpha*e is receiving work; discharge duration is W/mu. "
          "Linear mass discharge gives mean post-arrival time W/(2*mu). "
          "Hold M,p,L,lambda,alpha,mu fixed: de/dtau=-lambda*e and "
          "dtau/dv=-L/v^2 yield the displayed chain-rule derivative. Alpha is "
          "the total merge-work multiplier; alpha-1 is the extra part.",
          [VSL+"/extraction/s13_m2_source_to_equation_audit.md",
           VSL+"/synthesis/s13/s13_m2_certificate_analysis.tex",
           VSL+"/synthesis/s13/s13_mean_time_and_speed_derivative_explained.tex",
           "https://doi.org/10.1080/21680566.2019.1700846",
           "https://eprints.whiterose.ac.uk/id/eprint/77315/7/Carey%20et%20al%202013_with_coversheet.pdf",
           "https://doi.org/10.1080/23249935.2022.2033878"])
chain(s,["Common\nspeed","Residence\ntime","Remaining\nmerge mass","Receiving\nwork","Complete\ntime"],y=1.46)
formula(s,"vsl_cost_gradient",
        [r"J=\tau+\frac{M+(\alpha-1)e(\tau)}{2\mu}",
         r"\frac{dJ}{dv}=-\frac{L}{v^2}\left[1-\frac{(\alpha-1)\lambda e(\tau)}{2\mu}\right]"],
        .95,2.54,11.4,1.56,fontsize=27)
body(s,["J: total cohort time / M; all mass starts together and clears completely.",
        "M: cohort mass; e(τ): mass still in the ending lane at the choke.",
        "α>1: merge-work multiplier; μ: receiving work per unit time.",
        "Linear batch discharge: mean choke time = clearance duration / 2.",
        "Transfer: Nagalur Subraveti; receiving: Carey; common VSL: Zhang.",
        "Exponential transfer and batch reduction: analytical adaptations."],y=4.59,step=.43)

s = slide("VSL: an analytical task, with a study still ahead",
          "VSL S13 is properly closed as a conditional decision-relevant certificate. "
          "It is more than just admitting a model family: observer, action, dynamics, "
          "complete accounting and gradient relevance are specified for one analytical member. "
          "It remains incomplete as a model-family or empirical research study. "
          "Remaining mandatory-merge mass compresses the consequential information "
          "to a scalar; full lanes, order, gaps and vehicle identity have not been shown necessary. "
          "Robustness/causal deletion, minimum-state comparison, traffic-domain and "
          "reference admission, derivative verification and held-out independent usefulness "
          "are deferred. Further VSL review is set aside for colleague discussion.",
          [VSL+"/reporting/s13/s13_99_m2_certificate_sprint_closure.md",
           VSL+"/reporting/s13/s13_35_study_sufficiency_and_extension_audit.md",
           VSL+"/synthesis/s13/s13_deferred_research_readiness_review_plan.md"])
body(s,["Established: a conditional hidden-state → action → gradient chain.",
        "One continuous merge-preparation statistic may be sufficient.",
        "Not established: calibrated transfer law or real-world regime prevalence.",
        "Still needed: ablations, model-family robustness and stronger baselines.",
        "Then: a traffic-valid reference, derivative checks and independent usefulness.",
        "Further VSL review is deferred pending colleague discussion."],y=1.68,step=.72)

s = slide("Ramp metering: one local merge",
          "The branch is reopened method-first, not as a compatible implementation extension. "
          "One finite-storage on-ramp enters a two-mainline-lane local merge candidate. "
          "The common action is a bounded open-loop infrastructure release-rate schedule. "
          "Open-loop isolates model gradient correctness and one-step usefulness before "
          "a feedback/PID/MPC controller is introduced. Meter permission caps upstream "
          "release, not necessarily downstream instantaneous entry after propagation/storage. "
          "Complete road-total density/flow/speed/count history and ramp history, common "
          "known forcing/geometry/parameters, all-stream time through clearance are retained. "
          "Multiple ramps and detailed gap-timed/CAV control are not this first task.",
          [RAMP+"/reporting/rm7_20_method_first_task_formulation_and_gate.md",
           RAMP+"/synthesis/rm7_post_selection_design_inventory.md",
           RAMP+"/reporting/rm9_current_handoff.md"])
start=len(s.shapes)
for y in [1.84,2.62]:
    rect(s,.98,y,11.29,.55,ROAD)
    arrow(s,11.62,y+.16,.45)
rect(s,1.33,4.04,5.1,.55,ROAD)
line(s,6.43,4.31,8.9,2.90,end=True,width=2)
rect(s,3.48,3.93,.15,.77,INK)
rect(s,1.55,4.10,1.6,.43,AMBER)
rect(s,4.54,4.16,.55,.31,TEAL)
text(s,1.2,4.81,2.6,.42,"Finite ramp queue",role="plot",center=True)
text(s,3.26,3.23,1.13,.4,"Meter",role="plot",center=True,bold=True)
text(s,5.95,4.81,4.33,.42,"Release ≠ downstream entry",role="plot",center=True)
group(s,start,"local freeway ramp and stored releases")
body(s,["Local lane-aware CTM: longitudinal cells with finite storage and lane exchange.",
        "Open-loop rate changes isolate model sensitivity before controller design.",
        "Count ramp waiting, mainline time, external waiting and all residual traffic."],y=5.41,step=.52)

s = slide("Ramp CTM: lateral and forward flow",
          "The adopted kernel is a lane-aware cell-transmission model, with continuous "
          "vehicle-equivalent mass in each physical mainline lane and longitudinal cell, "
          "and discrete time steps. It is not one lumped mass for an entire lane or "
          "individual vehicle dynamics. Three adjacent cells are shown schematically, "
          "not as an approved cell count or completed network geometry. First apply "
          "bounded paired density-driven lateral transfers within every cell; these "
          "conserve its two-lane total. Then recompute ordinary CTM sending and receiving "
          "from post-lateral masses. At an ordinary interface, accepted lane flow is "
          "the minimum of upstream sending and downstream receiving. The time-step "
          "mass balance adds h times total accepted inflow f and subtracts h times "
          "outflow q. The ordinary min formula is not applied independently to each "
          "input at a merge: mainline and ramp still share one receiver under positive "
          "input priorities with unused-supply redistribution. Metered ramp releases "
          "must propagate through storage before entering that junction; exact "
          "ramp geometry and boundary transcription remain open. Transverse and "
          "longitudinal resources are separate in the approved adaptation: no "
          "original-budget-minus-lateral cap remains. Lane exchange changes occupancy "
          "and therefore flow offers and storage availability, not physical lane "
          "capacity; throughput improvement is not guaranteed. Nagalur Subraveti "
          "supplies lane-wise CTM/exchange antecedents and Wright et al. supplies "
          "the node allocator; their composition is an approved review adaptation, "
          "not a reproduced calibrated paper. The FD functions and accepted "
          "time-step/storage bounds are in the source audit. Diagram blocks are "
          "symbolic occupancy, not computed mass or traffic results. The visible lateral "
          "equations now retain the accepted bounded operator: A is a directed "
          "preference between zero and one, positive only when the density advantage "
          "passes the approved threshold; lambda is limited by that fraction of "
          "the donor pre-exchange sending and by recipient pre-exchange receiving "
          "in the same longitudinal cell. Multiplying the incoming-minus-outgoing "
          "lateral rate by h gives each lane's post-exchange mass. The terms cancel "
          "when summing lanes. This D/S-bounded exchange inherits the approved "
          "time-step/storage restrictions. A is not a new control action; the "
          "open-loop meter rate remains the common actuator. The exact normalized "
          "density-contrast preference, its propensity/threshold parameters and "
          "empty-cell convention remain defined by the existing source audit. "
          "Forward D/S are recomputed after the lateral stage, as displayed.",
          [RAMP+"/discussion/rm9_resource_b_approval_and_continuation.md",
           RAMP+"/synthesis/rm9/rm9_separate_resource_and_observer_audit.tex",
           "https://repository.tudelft.nl/file/File_2c4359e7-5e5a-4651-964f-6f3610ab89ae",
           "https://horowitz.me.berkeley.edu/Publications_files/All_papers_numbered/Wright_node_models_transportation_networks_Transportation_research_B_2017.pdf"])
start=len(s.shapes)
text(s,.51,1.83,1.06,.4,"Lane 1",role="plot",center=True)
text(s,.51,2.43,1.06,.4,"Lane 2",role="plot",center=True)
for x,label in [(1.72,"Cell i − 1"),(5.16,"Cell i"),(8.60,"Cell i + 1")]:
    text(s,x,1.39,3.0,.4,label,role="plot",center=True,bold=True,color=TEAL)
    for y in [1.82,2.42]:
        rect(s,x,y,3.0,.46,ROAD)
    line(s,x,1.80,x,2.91,color=GRAY,width=1)
    line(s,x+3.0,1.80,x+3.0,2.91,color=GRAY,width=1)
for x,top,bottom in [(1.88,.75,1.50),(5.32,.90,1.55),(8.76,1.25,.85)]:
    rect(s,x,1.90,top,.30,TEAL)
    rect(s,x,2.50,bottom,.30,AMBER)
arrow(s,7.57,2.12,.30,.51,color=AMBER,kind=MSO_SHAPE.UP_ARROW)
for x in [4.79,8.23,11.70]:
    for y in [1.94,2.54]:
        arrow(s,x,y,.33,color=TEAL)
text(s,.9,3.01,5.45,.39,"Within cell: paired lateral exchange",role="plot",center=True)
text(s,6.50,3.01,5.65,.39,"Across cells: forward transmission",role="plot",center=True)
text(s,.73,3.85,1.42,.4,"Lateral",role="plot",bold=True,color=AMBER)
text(s,.73,4.94,1.42,.4,"Forward",role="plot",bold=True,color=TEAL)
group(s,start,"CTM cells and two-stage lateral-forward update")
formula(s,"ramp_resource",
        [r"\lambda_{i,\ell j}=\min\!\left[A_{i,\ell j}D_{i\ell}(n_{i\ell}),\;S_{ij}(n_{ij})\right]",
         r"n_{i\ell}^{\mathrm{after}}=n_{i\ell}+h\left(\lambda_{i,j\ell}-\lambda_{i,\ell j}\right)",
         r"q_{i\ell}=\min\!\left[D_{i\ell}(n_{i\ell}^{\mathrm{after}}),\;S_{i+1,\ell}(n_{i+1,\ell}^{\mathrm{after}})\right]",
         r"n_{i\ell}^{\mathrm{next}}=n_{i\ell}^{\mathrm{after}}+h\left(f_{i\ell}-q_{i\ell}\right)"],
        2.15,3.49,10.0,2.13,fontsize=25,line_spacing=.66)
body(s,["n: lane-cell mass; i: cell; ℓ, j: lanes; h: time step; after: post-exchange.",
        "A: bounded density preference; λ: directed lateral rate; D/S: sending/receiving.",
        "q: forward rate; f: accepted inflow; ramp + mainline share receiving.",
        "Lane CTM/exchange: Nagalur Subraveti; node priorities: Wright. Adapted."],y=5.79,step=.40)

s = slide("Ramp lane contrast is not yet certified hidden",
          "Lane contrast is a candidate state omission, not an established hidden-information "
          "result. On the strict ordinary free-flow B branch, total post-lateral mass and "
          "aggregate actual flow reconstruct the two-lane split when known mobilities differ. "
          "A separate passage-speed observation is retained, not replaced by Q/K. "
          "The geometric illustration uses normalized c1=1,c2=1.5, masses (2,8) versus "
          "(6,4), both total 10 but flows 14 and 12. This is hand algebra, not matched "
          "history or calibrated traffic. A pre-transfer map can also collapse states; "
          "if all remaining state/forcing agree and actions act only afterwards, that "
          "lost coordinate cannot change actions. Congested/receiving-limited states "
          "need the full observation map and complete prior-history audit.",
          [RAMP+"/synthesis/rm9/rm9_separate_resource_and_observer_audit.tex",
           RAMP+"/extraction/rm9_separate_resource_claims.md",
           RAMP+"/reporting/rm9_50_separate_resource_findings_and_observer_gate.md"])
start=len(s.shapes)
text(s,2.03,1.38,9.15,.39,"Free-flow illustration: normalized lane mobilities 1 and 1.5.",
     role="plot",center=True)
for x, masses, flow in [(1.0,[2,8],14),(7.0,[6,4],12)]:
    for j,(mass,col) in enumerate(zip(masses,[TEAL,AMBER])):
        y=1.94+.78*j
        rect(s,x,y,4.9,.53,ROAD)
        rect(s,x,y,mass/10*4.9,.53,col)
        text(s,x+5.0,y+.04,.66,.4,str(mass),role="plot",bold=True)
    text(s,x,3.43,5.2,.45,f"Total mass 10; forward flow {flow}",role="plot",center=True,bold=True)
text(s,2.65,4.11,8.45,.40,"Different lane mobilities make aggregate flows distinguishable.",role="plot",center=True)
group(s,start,"equal mass does not mean equal aggregate information")
body(s,["Candidate omission: target-lane occupancy and local redistribution.",
        "Free-flow density and flow can already reveal the lane split.",
        "Congestion requires a precise independent speed-and-history map.",
        "No equal-history pair with different useful metering actions is established."],y=5.03,step=.49)

s = slide("Ramp gradients follow final departure timing",
          "The conceptual chain is common meter parameter theta -> releases -> post-meter "
          "storage/propagation -> shared receiving -> accepted merge and lane response -> "
          "complete final departures -> all-stream system time. The displayed sensitivity "
          "identity assumes fixed common arrivals/initial mass and clearance horizon, "
          "complete population accounting and fixed active branches. E is final system "
          "exit flow, not local ramp merge entry; P is all waiting/road/ramp mass. "
          "Blocked meter perturbations may have zero derivative. Queue relocation with "
          "identical departures cannot improve total time. Smooth active-branch composition "
          "is available but whole-chain numerical correctness and independent usefulness "
          "remain unestablished; no numerical traffic or gradient tests are claimed.",
          [RAMP+"/synthesis/rm9/rm9_separate_resource_and_observer_audit.tex",
           RAMP+"/synthesis/rm8_task_gradient_reference_screen.md"])
chain(s,["Meter\nrate","Release +\nstorage","Accepted\nmerging","Final\ndepartures","All-stream\ntime"],y=1.45)
formula(s,"ramp_objective_gradient",
        [r"J=\int_0^T P(t)\,dt",
         r"\frac{\partial J}{\partial\theta}=-\int_0^T(T-t)\,\frac{\partial E(t)}{\partial\theta}\,dt"],
        .85,2.54,11.6,1.85,fontsize=26,line_spacing=1.0)
body(s,["J: total system time; P: all traffic still in system; E: final exit flow.",
        "T: common clearance horizon; θ: open-loop meter-schedule parameter.",
        "Moving waiting from mainline to ramp alone does not reduce total time.",
        "Piecewise sensitivity is formulated; correctness and usefulness remain open."],y=4.83,step=.52)

s = slide("Different evidence levels; one common test",
          "Synthesis as of 2 October 2026: VSL S13 closed positively only as a "
          "conditional analytical member. Ramp RM9 is active at the G1 observer-contract "
          "gate; UG3.B resource amendment is approved, observer alternatives UG4.A/B "
          "are not. Idealized FD passage-speed benchmark is recommended but not adopted; "
          "measurement/source-first review is the other route. No complete ramp "
          "history/action witness or independently useful gradient exists. The next "
          "measurement decision may alter identifiability and therefore cannot be assumed. "
          "Both tasks allow aggregate-plus-scalar sufficiency rather than require full "
          "lane resolution, vehicle identity or discreteness. This slide deck does not "
          "advance any research gate or authorize experiments.",
          [VSL+"/reporting/s13/s13_99_m2_certificate_sprint_closure.md",
           RAMP+"/reporting/rm9_current_handoff.md",
           RAMP+"/reporting/rm9_50_separate_resource_findings_and_observer_gate.md"])
body(s,["VSL: conditional proof of hidden information changing action and gradient.",
        "Ramp metering: conservative formulation; information value remains open.",
        "Both: a continuous scalar may suffice; full lane detail is not presumed.",
        "Next ramp question: what does speed measure, and when is it observed?",
        "An idealized benchmark or a measurement-first review remains to be chosen.",
        "Research-ready claims need robustness and independent outcome testing."],
     y=1.67,step=.73)


def leaves(shapes, ox=0, oy=0, sx=1, sy=1):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xf = shape._element.grpSpPr.xfrm
            co, ce = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
            gx, gy = shape.width/int(ce.get("cx")), shape.height/int(ce.get("cy"))
            yield from leaves(shape.shapes,
                              ox+sx*(shape.left-int(co.get("x"))*gx),
                              oy+sy*(shape.top-int(co.get("y"))*gy),sx*gx,sy*gy)
        else:
            yield shape, ((ox+sx*shape.left)/EMU,(oy+sy*shape.top)/EMU,
                          sx*shape.width/EMU,sy*shape.height/EMU)


def font(points,bold=False,scale=144):
    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans"
                              +("-Bold" if bold else "")+".ttf",round(points*scale/72))


def validate(saved):
    assert len(saved.slides)==14 and sum(r["boilerplate"] for r in records)==2
    counts=[]; geometric=[]; visible=[]
    for i,s in enumerate(saved.slides):
        lines=0; equations=0; titles=0
        assert s.background.fill.fore_color.rgb==rgb(WHITE)
        geometric.append(sum(sh.shape_type==MSO_SHAPE_TYPE.GROUP for sh in s.shapes))
        for sh,(x,y,w,h) in leaves(s.shapes):
            assert min(x,y)>=-.01 and x+w<=13.344 and y+h<=7.51,(i+1,sh.name,"off slide")
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE and sh.name.startswith("math:"):
                equations+=math_lines[sh.name]
            if not sh.has_text_frame or not sh.text:
                continue
            visible.append(sh.text)
            istitle=sh.name.startswith("title:")
            titles+=int(istitle)
            if not istitle and not sh.name.startswith("plot:"):
                lines+=len(sh.text.splitlines())
            tf=sh.text_frame
            for p in tf.paragraphs:
                pts=32 if istitle else 18
                assert p.font.name==FONT and p.font.size.pt==pts,(i+1,sh.name,"font")
                avail=(w-(tf.margin_left+tf.margin_right)/EMU)*144
                assert font(pts,bool(p.font.bold)).getlength(p.text)<=avail,(i+1,sh.name,"text width")
            required=len(sh.text.splitlines())*(32 if istitle else 18)*1.15/72
            availh=h-(tf.margin_top+tf.margin_bottom)/EMU
            assert required<=availh+.04,(i+1,sh.name,"text height")
        assert titles==1 and lines<=8 and lines+equations<=8,(i+1,lines,equations)
        assert len(s.notes_slide.notes_text_frame.text)>250
        counts.append({"slide":i+1,"body_lines":lines,"formula_lines":equations})
    assert not re.search(r"\b(S0[1-9]|S1[0-3]|RM[0-9]|UG[1-4]|Candidate [ABC]|M2C|Phase [A-F])\b","\n".join(visible))
    for record in records:
        for source in record["sources"]:
            if "://" not in source:
                assert (ROOT/source).exists(),source
    with zipfile.ZipFile(OUT) as archive:
        assert archive.testzip() is None
        names=set(archive.namelist())
        for member in names:
            if member.endswith(".xml") or member.endswith(".rels") or member.endswith(".svg"):
                etree.fromstring(archive.read(member))
        vector_count=sum(name.endswith(".svg") for name in names)
        assert vector_count==5
        svg_refs=0
        for i in range(1,15):
            xml=etree.fromstring(archive.read(f"ppt/slides/slide{i}.xml"))
            rels=etree.fromstring(archive.read(f"ppt/slides/_rels/slide{i}.xml.rels"))
            relmap={node.get("Id"):node.get("Target") for node in rels}
            for node in xml.findall(".//{"+SVG_NS+"}svgBlip"):
                rid=node.get(qn("r:embed"))
                assert rid in relmap
                target=posixpath.normpath(posixpath.join("ppt/slides",relmap[rid]))
                assert target in names and target.endswith(".svg")
                svg_refs+=1
        assert svg_refs==5
    return {"slides":14,"boilerplate":2,"substantive":12,"body_line_counts":counts,
            "grouped_diagrams_by_slide":geometric,"embedded_vector_assets":vector_count,
            "font":"Calibri","title_points":32,"body_points":18,
            "white_background":True,"xml_zip_and_relationship_checks":"PASS",
            "source_paths":"PASS","layout_checks":"PASS",
            "preview":"Pillow layout proof using DejaVu Sans; not Office rendering"}


def color_of(value,default):
    try:return "#"+str(value.rgb)
    except (AttributeError,ValueError,TypeError):return default


def preview(saved):
    scale=120; pages=[]
    for i,s in enumerate(saved.slides,1):
        im=Image.new("RGB",(1600,900),"#"+WHITE)
        draw=ImageDraw.Draw(im)
        for sh,(x,y,w,h) in leaves(s.shapes):
            x,y,w,h=(v*scale for v in (x,y,w,h))
            bounds=(x,y,x+w,y+h)
            if sh.shape_type==MSO_SHAPE_TYPE.LINE:
                xf=sh._element.spPr.xfrm
                xx=(x+w,x) if xf.get("flipH")=="1" else (x,x+w)
                yy=(y+h,y) if xf.get("flipV")=="1" else (y,y+h)
                col=color_of(sh.line.color,"#"+GRAY)
                draw.line((xx[0],yy[0],xx[1],yy[1]),fill=col,width=3)
                if sh._element.spPr.ln.find(qn("a:tailEnd")) is not None:
                    angle=math.atan2(yy[1]-yy[0],xx[1]-xx[0]);tip=(xx[1],yy[1])
                    draw.polygon([tip,(tip[0]-13*math.cos(angle-.35),tip[1]-13*math.sin(angle-.35)),
                                  (tip[0]-13*math.cos(angle+.35),tip[1]-13*math.sin(angle+.35))],fill=col)
                continue
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                picture=Image.open(BytesIO(sh.image.blob)).convert("RGBA")
                picture=picture.resize((max(1,round(w)),max(1,round(h))),Image.Resampling.LANCZOS)
                im.paste(picture,(round(x),round(y)),picture)
                continue
            if sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill=color_of(sh.fill.fore_color,None) if sh.fill.type==1 else None
                col=color_of(sh.line.color,None)
                kind=sh.auto_shape_type
                if kind==MSO_SHAPE.OVAL:
                    draw.ellipse(bounds,fill=fill,outline=col,width=2)
                elif kind in [MSO_SHAPE.RIGHT_ARROW,MSO_SHAPE.UP_ARROW]:
                    if kind==MSO_SHAPE.RIGHT_ARROW:
                        points=[(x,y+h*.25),(x+w*.62,y+h*.25),(x+w*.62,y),
                                (x+w,y+h*.5),(x+w*.62,y+h),(x+w*.62,y+h*.75),(x,y+h*.75)]
                    else:
                        points=[(x+w*.25,y+h),(x+w*.25,y+h*.38),(x,y+h*.38),
                                (x+w*.5,y),(x+w,y+h*.38),(x+w*.75,y+h*.38),(x+w*.75,y+h)]
                    draw.polygon(points,fill=fill,outline=col)
                else:
                    draw.rectangle(bounds,fill=fill,outline=col,width=2)
            if not sh.has_text_frame or not sh.text:continue
            tf=sh.text_frame
            ml,mt=tf.margin_left/EMU*scale,tf.margin_top/EMU*scale
            mr,mb=tf.margin_right/EMU*scale,tf.margin_bottom/EMU*scale
            paragraphs=list(tf.paragraphs);pts=paragraphs[0].font.size.pt
            step=pts*scale/72*1.15;ty=y+mt
            if tf.vertical_anchor==MSO_ANCHOR.MIDDLE:
                ty=y+mt+(h-mt-mb-len(paragraphs)*step)/2
            for p in paragraphs:
                f=font(p.font.size.pt,bool(p.font.bold),scale)
                tx=x+ml
                if p.alignment==PP_ALIGN.CENTER:
                    tx+=(w-ml-mr-f.getlength(p.text))/2
                draw.text((tx,ty),p.text,font=f,fill=color_of(p.font.color,"#"+INK),anchor="lt")
                ty+=step
        im.save(ASSETS/f"slide_{i:02d}_layout.png")
        pages.append(im)
    contact=Image.new("RGB",(1366,7*401+8),"#DCE1E5")
    for i,im in enumerate(pages):
        thumb=im.resize((667,375),Image.Resampling.LANCZOS)
        contact.paste(thumb,(8+(i%2)*683,8+(i//2)*401))
    contact.save(PREVIEW)
    for start in [0,8]:
        subset=pages[start:start+8]; rows=math.ceil(len(subset)/2)
        sheet=Image.new("RGB",(1366,rows*401+8),"#DCE1E5")
        for j,im in enumerate(subset):
            sheet.paste(im.resize((667,375),Image.Resampling.LANCZOS),
                        (8+(j%2)*683,8+(j//2)*401))
        sheet.save(ASSETS/f"contact_{start+1:02d}_{start+len(subset):02d}.png")


def write_notes(result):
    content=["# Weekly freeway-control review deck — request 6",
             "",
             "Date: 2026-10-02. Colleague-facing methodology/findings, from VSL task "
             "definition to the current ramp-metering review. Two boilerplate and "
             "12 substantive slides; white background, Calibri 32/18 pt.",
             "",
             "## Content and claim boundaries",
             "",
             "Objective clarification: slide 8 uses J for VSL mean cohort time "
             "(total time divided by fixed mass); slide 13 uses J for ramp total "
             "system time. Their labels now make the normalization explicit. "
             "The [three-page explanation](../differentiable_freeway_bottleneck_task_review/synthesis/s13/s13_mean_time_and_speed_derivative_explained.pdf) "
             "defines all VSL symbols and derives the linear-batch cost and "
             "speed gradient. Its editable LaTeX source is retained alongside it. "
             "These are explanatory corrections, not changed model semantics.",
             "",
             "VSL is complete as a conditional analytical task/certificate, not an "
             "empirical or full model-family study. The closed proof preserves one "
             "common lane-uniform action. Ending-lane fraction can alter the initial "
             "optimal action and gradient under its stated inherited history collision "
             "and transfer/receiving regime. Its exponential transfer and batch service "
             "are analytical adaptations, not a calibrated law from one paper.",
             "",
             "Ramp metering is active at the observer-contract decision. The separate "
             "resource kernel is approved for analytic review; the idealized speed "
             "map remains proposed. No complete ramp equal-history/action witness "
             "or independently useful gradient is claimed. Gate labels and sprint "
             "chronology are kept out of visible slides.",
             "",
             "The normalized VSL plot evaluates the existing certificate's closed-form "
             "illustration only. It is not a new traffic simulation, calibration or "
             "parameter experiment. All diagrams are grouped native shapes or one "
             "vector picture object. Sources and detailed qualifications are in "
             "each slide's PowerPoint speaker notes.",
             "",
             "## Build and format",
             "",
             "Generator: [generate_weekly_freeway_control_review_20261002.py](generate_weekly_freeway_control_review_20261002.py)",
             "",
             "Build command:",
             "",
             "    /home/zpz/miniconda3/envs/differential_sim/bin/python misc/slides/generate_weekly_freeway_control_review_20261002.py",
             "",
             "Dependencies already installed: python-pptx 1.0.2, matplotlib 3.10.9, "
             "Pillow, lxml and NumPy. No installation or production changes.",
             "",
             "PowerPoint text is editable Calibri. Mathematical expressions and the "
             "analytic plot use scalable SVG with a PNG fallback, following "
             "[Microsoft's SVG picture convention](https://learn.microsoft.com/en-us/openspecs/office_standards/ms-odrawxml/7e1f1524-1569-4aa2-a6c9-aab2d855bd48). "
             "Equation sources remain in the generator and SVGs in the asset folder. "
             "No complex plain-text equation is used.",
             "",
             "## Validation and rendering limit",
             "",
             "Checks cover slide/boilerplate counts, Calibri 32/18 metadata, body lines "
             "(also conservatively counting equation lines), grouped diagrams, text "
             "width/height and slide bounds, source paths, ZIP/XML and embedded "
             "SVG relationship targets. All pass; see "
             "[validation JSON](weekly_freeway_control_review_20261002_validation.json).",
             "",
             "Pillow layout proof reads the saved PPTX, including group transforms "
             "and PNG fallbacks. DejaVu Sans substitutes for locally unavailable "
             "Calibri; no Office/LibreOffice renderer is installed. This is an "
             "inspected layout proof, not a native PowerPoint-rendering claim.",
             "",
             "Manual preview inspection checked all 14 slides, plus full-size "
             "transfer and gradient slides. Labels/titles were wrapped or shortened "
             "to fit without reducing the requested fonts; the two integral rows "
             "were spaced apart after inspection. No scientific simulation or "
             "numerical gradient test was required or performed. VSL remains "
             "closed at its conditional analytical scope; the ramp observer gate "
             "remains open. Next allowed review action is discussion of that gate.",
             "",
             "## Slide-by-slide speaker notes",
             ""]
    for record in records:
        content += [f"### {record['slide']}. {record['title']}", "", record["notes"], "", "Sources:", ""]
        for source in record["sources"]:
            if "://" in source:
                content.append("- [Primary source]("+source+")")
            else:
                rel=os.path.relpath(ROOT/source,HERE)
                content.append("- ["+source+"]("+rel+")")
        content.append("")
    NOTES.write_text("\n".join(content)+"\n")


if __name__=="__main__":
    prs.save(OUT)
    saved=Presentation(OUT)
    result=validate(saved)
    preview(saved)
    write_notes(result)
    REPORT.write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))
    print(OUT)
    print(PREVIEW)
