"""Historical freeway-review synthesis: two boilerplate + five content slides.

Reuse the existing deck's native PowerPoint styling and layout renderer without
changing or rebuilding that deck. All scientific content ends at the pre-review
handover of 16 September 2026.
"""

from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from types import SimpleNamespace
import math
import re
import zipfile

from PIL import Image
from pptx import Presentation


HERE = Path(__file__).resolve().parent
OUT = HERE / "freeway_review_to_task_first_handoff_20260918.pptx"
PREVIEW = HERE / "freeway_review_to_task_first_handoff_20260918_preview.png"
spec = spec_from_file_location("freeway_style", HERE / "generate_freeway_review_deck_20260918.py")
style = module_from_spec(spec)
spec.loader.exec_module(style)
style.prs = Presentation()
style.prs.slide_width = style.Inches(13.333333)
style.prs.slide_height = style.Inches(7.5)
prs = style.prs
prs.core_properties.title = "From model compatibility to task-first freeway research"
prs.core_properties.subject = "Methodology and findings through the 16 September 2026 handover"
prs.core_properties.author = "Research review"
t, rect, node = style.text, style.rect, style.node
arrow, line, group = style.arrow, style.line, style.group_since


def slide(title, notes):
    return style.new_slide(title, notes)


def body(s, y, value, **kwargs):
    return t(s, .85, y, 11.65, .42, value, **kwargs)


# Exactly two boilerplate pages, reinstated by request 5's default rules.
s = slide("From model-first to task-first research", """Title slide. Scope: the original ramp-metering feasibility review opened on 14 September 2026, its bounded conclusion and closure on 15 September, and the task-first correction and pre-review handover on 16 September. This is a historical methodology synthesis. Subsequent N1/N2/Route B and analytical/reference-bridge findings are outside this deck's requested scope, even though the reference style deck discusses those later records.
Sources: ../selective_discreteness_ramp_metering_review/README.md; reporting/closure_record.md in that package; ../differentiable_freeway_bottleneck_task_pre_review/pre_review_closure_and_handoff.md.
""")
t(s, .85, 2.55, 11.65, .55, "Freeway control: what traffic information is actually necessary?", bold=True, color=style.TEAL)
t(s, .85, 3.50, 11.65, .45, "Methodological findings through 16 September 2026")
t(s, .85, 5.90, 11.65, .42, "Technical discussion for colleagues")

s = slide("Contents", """Contents slide. Five substantive slides organize the argument semantically, with no sprint, phase, gate, source-access or progress reporting in the visible content. Short source descriptors appear where they help convey the finding; full source boundaries and provenance remain in the speaker notes and companion work note. No simulation, derivative result or minimum-state necessity claim is introduced.
The argument moves from the original joint compatibility/value criterion, to evidence and its limits, to the bounded transfer case, to the corrected task-first hierarchy, and finally to the operational anchor and unresolved scientific question inherited by the handover.
""")
for y, value in zip([1.70, 2.60, 3.50, 4.40, 5.30], [
    "1   The original compatibility and scientific-value test",
    "2   Strong aggregate alternatives and ambiguous mesoscopic value",
    "3   Why the ordered-stream case is a transfer test",
    "4   The correction: task before model compatibility",
    "5   Lane reduction as an independent research anchor",
]):
    body(s, y, value)

s = slide("The original question imposed two conditions", """The original review asked whether ramp metering could become an orthogonal parallel task within an already established architecture. It did not forbid negative or macro-sufficiency results, but strict compatibility was an initial admissibility condition. The two circles are conceptual sets, not quantitative findings: one is tasks compatible with the current continuous-mass first-order architecture, the other tasks independently diagnostic of selected mesoscopic information. The overlap is the desired role, not an established candidate.
Fixed architecture in that review: aggregate first-order propagation, continuous traffic mass, fixed structural routes/topology, and localized ordered metadata. Compatibility alone was insufficient; scientific diagnosticity and practical evidence feasibility were separate judgments. The original evidence audit did not support promotion of any reviewed formulation meeting both conditions. No universal nonexistence claim follows.
Sources: ../selective_discreteness_ramp_metering_review/review_protocol_plan.md, sections 1, 2 and 6; reporting/colleague_brief/ramp_metering_parallel_task_assessment.tex, sections 1 and 6; reporting/closure_record.md.
""")
body(s, 1.40, "Starting point: continuous-mass, first-order physics with localized route order.")
start = len(s.shapes)
rect(s, 2.00, 2.32, 5.30, 3.35, None, style.BLUE, kind=style.MSO_SHAPE.OVAL)
rect(s, 6.00, 2.32, 5.30, 3.35, None, style.TEAL, kind=style.MSO_SHAPE.OVAL)
t(s, 2.45, 3.25, 3.80, .42, "Compatible with", center=True, bold=True, color=style.BLUE)
t(s, 2.45, 3.76, 3.80, .42, "the existing model", center=True)
t(s, 7.10, 3.25, 3.70, .42, "Distinct information", center=True, bold=True, color=style.TEAL)
t(s, 7.10, 3.76, 3.70, .42, "needed for control", center=True)
t(s, 5.67, 4.50, 1.98, .42, "Both needed", center=True, bold=True)
group(s, start, "Joint compatibility and independent scientific-value conditions")
body(s, 6.20, "No reviewed formulation established both conditions for a parallel task.", bold=True)
body(s, 6.76, "This is a conclusion within the original scope; it does not reject freeway control.")

s = slide("Aggregate alternatives already set a high bar", """This slide combines the earlier evidence mapping and debate into the scientific finding, rather than reporting their procedural sequence. ALINEA and CTM-based optimal/distributed control establish strong controller and lower-resolution alternatives. Reilly et al. (2015), RM0024, provide direct coordinated I-15 ramp-metering gradients via a manually derived discrete adjoint. Gu et al. (2022), RM0026, supply abstract/metadata-level evidence for smoothed CTM ramp-metering MPC; no detailed equation claim is made. Goatin et al. (2026), RM0025, provide full-text multiclass network adjoint evidence, not a completed version of the proposed ordered-stream task. Poole and Kotsialos (2016), RM0030, concern freeway METANET calibration, not ramp-control performance. These derivative constructions have different forward/backward contracts and are not interchangeable.
The closed ramp review found mesoscopic work emphasizing routes/DTA, event progression, stochastic realization and scale, with some lane refinement. The pre-review terminology scan also identified packet/cohort, class/lane, distributional and hybrid local-interaction families. This is semantic orientation, not a fidelity ladder or proof of minimum resolution. Exact local FIFO rank was not established as necessary by the reviewed ramp corpus; absence in this bounded corpus is not universal absence.
Sources: ../selective_discreteness_ramp_metering_review/synthesis/feedback20_evidence_debate.md, sections 1-3 and 5; archive/metadata/feedback20_sources.md; ../differentiable_freeway_bottleneck_task_pre_review/freeway_mesoscopic_meaning_quick_scan.md, sections 1-3 and 6.
Primary anchors in existing records: https://doi.org/10.1007/s10957-015-0749-1 ; https://doi.org/10.1109/TITS.2021.3123815 ; https://doi.org/10.1007/s11081-026-10122-4 ; https://doi.org/10.1016/j.trc.2016.07.008 . No new literature search was performed for this deck.
""")
rows = [
    ("Aggregate control is a strong starting point.", True, style.BLUE),
    ("Feedback / optimal control: ALINEA and CTM-based ramp control.", False, style.INK),
    ("Efficient sensitivities: Reilly's adjoint; Gu's smoothed CTM.", False, style.INK),
    ("Multiclass adjoint: Goatin; freeway AD calibration: Poole.", False, style.INK),
    ("Mesoscopic value varies with retained information and behavior.", True, style.TEAL),
    ("Routes, cohorts, classes, lanes, distributions and local interactions.", False, style.INK),
    ("The reviewed evidence does not establish a need for exact queue order.", False, style.INK),
    ("Differentiability alone does not establish useful control gradients.", True, style.TEAL),
]
for y, (value, bold, color) in zip([1.52, 2.12, 2.72, 3.32, 4.26, 4.86, 5.46, 6.37], rows):
    body(s, y, value, bold=bold, color=color)

s = slide("The ordered-stream case tests transfer", """Conceptual diagram, not simulated data. A metered upstream stream enters a first-order link, route-labelled order propagates downstream, and a downstream exit restriction creates order-sensitive service. Blue packet rectangles denote exit-bound mass and teal rectangles continuing mass; rectangle counts and lengths carry no quantitative meaning. The amber branch denotes constrained exit service. This case was coherent only with non-overlapping merge-input eligibility intervals and a restricted point geometry, avoiding an unsupported unique chronology of simultaneous continuous arrivals.
The scientific finding was partial causal-chain novelty, but no orthogonal representation hypothesis: the decisive mechanism is the same route-labelled ordered access tested by the signal-control application. It could remain an optional propagation/transfer or ablation fixture; it was neither a selected parallel task nor an implementation obligation. More realistic overlapping and multilane cases would require additional chronology/lateral behavior. A simple aggregate ramp case could also remain an optional correctness fixture.
Sources: ../selective_discreteness_ramp_metering_review/review_protocol_plan.md, opening closure summary; synthesis/feedback20_evidence_debate.md, section 4; reporting/colleague_brief/ramp_metering_parallel_task_assessment.tex, sections 4-5; reporting/closure_record.md, Final disposition.
""")
start = len(s.shapes)
rect(s, 1.00, 2.55, 10.35, .66, style.ROAD)
line(s, 1.70, 4.18, 3.24, 2.88, width=5, color=style.GRAY, end=True)
rect(s, 2.20, 3.49, .15, .52, style.AMBER)
line(s, 9.70, 2.89, 11.18, 4.15, width=4, color=style.AMBER, end=True)
arrow(s, 10.83, 2.77, color=style.GRAY)
for x, color in [(4.35, style.BLUE), (4.90, style.TEAL), (5.45, style.BLUE), (6.00, style.TEAL)]:
    rect(s, x, 2.73, .38, .28, color)
t(s, 1.02, 4.36, 2.42, .42, "Upstream meter", center=True)
t(s, 4.38, 1.92, 3.60, .42, "First-order propagation", center=True)
t(s, 9.46, 4.36, 3.00, .42, "Restricted exit", center=True)
group(s, start, "Metered stream, ordered route mass and remote exit restriction")
body(s, 5.12, "Route-labelled order reaches a downstream service restriction.")
body(s, 5.57, "Non-overlapping inputs avoid an unresolved merge chronology.")
body(s, 6.02, "The decisive ordered-access mechanism is shared with the signal task.")
body(s, 6.47, "Useful as an optional transfer test; not an independent second hypothesis.", bold=True, color=style.TEAL)
body(s, 6.92, "Ordinary overlapping or multilane traffic needs additional modeling.")

s = slide("A scope rejection is not a task rejection", """The colleague exchange exposed the distinction between ruling out a task inside an existing model and evaluating the task independently. Philipp and Wang Yu asked to define tasks and questions first, then decide which model components are necessary. The earlier bounded ramp-review conclusion remains valid under its compatibility-framed question. The new review changes the admissibility hierarchy; it does not overwrite or universalize that finding.
Top row: the old fixed-model compatibility question. Bottom row: the task-first correction. Compatibility/reuse, engineering cost and paper organization are later audits after task validity and minimum-model justification. Acceleration, safety gaps or lane changes absent from the current project are possible task requirements, not automatic evidence against a task; they do impose scientific isolation, calibration and validation costs. The common ground is the question of necessary traffic resolution and the evidence logic, not a shared simulator, mesoscopic definition, codebase or paper. All resolutions and negative/infeasible outcomes are admissible.
Sources: ../../temp_content/feedback/email_trace_20260916.md, messages 2 and 4; ../differentiable_freeway_bottleneck_task_pre_review/feedback22_email_trace_response.md, opening clarification and sections 2, 6.1, 6.7; pre_review_closure_and_handoff.md, principles P1-P8.
""")
body(s, 1.43, "The earlier rejection answered a question bounded by model compatibility.")
start = len(s.shapes)
for x, label in [(1.0, "Fixed model"), (5.0, "Compatible task"), (9.0, "Value within scope")]:
    node(s, x, 2.23, 3.30, label, color=style.GRAY)
arrow(s, 4.45, 2.59); arrow(s, 8.45, 2.59)
for x, label in [(1.0, "Operational task"), (5.0, "Strong baseline"), (9.0, "Minimum state")]:
    node(s, x, 4.44, 3.30, label)
arrow(s, 4.45, 4.80, color=style.TEAL); arrow(s, 8.45, 4.80, color=style.TEAL)
group(s, start, "Old compatibility hierarchy and corrected task-first hierarchy")
body(s, 6.29, "Shared resolution question; separate models and papers remain admissible.", bold=True, color=style.TEAL)

s = slide("Lane reduction anchors the new question", """This slide states the inherited task-framing direction, not later formal-review findings. Temporary or permanent freeway lane-capacity reduction is the operational anchor; closure location, affected lanes and closure/reopening schedule are exogenous in the first formulation. Low, near-capacity and high-demand regimes are relevant scenario dimensions. Reopening does not guarantee that an existing queue has cleared. The road drawing is a conceptual lane reduction, not a selected three-to-two reference geometry or simulated trajectory.
F1 speed harmonization/advisory control and F2 coordinated ramp metering are co-primary initial candidates. A human speed recommendation is not a direct CAV command: penetration, compliance, response delay and heterogeneity must be specified in any eventual task. F4 recovery is a required scenario dimension if temporary closure is selected. F3 joint control and F5 cooperative lane-change/merge control are expansions or boundary/separate-project candidates. They are not preselected control tasks. Objective accounting must include the same traveler population and upstream/ramp waiting; fairness, ramp storage, speed variance and safety proxies require explicit reporting or constraints where appropriate.
Research logic: operational task; strongest low-resolution baseline; controlled information-loss/decision gap; minimum additional state; derivative correctness; useful action directions or rankings in an independent richer evaluator; cost and evidence feasibility. Model detail and task selection are open at this handover. No N1 result, minimum-state result, experiment or simulator development is asserted or authorized.
Sources: ../differentiable_freeway_bottleneck_task_pre_review/pre_review_closure_and_handoff.md, sections 2-5 and 8; feedback22_email_trace_response.md, task candidates and sections 6.2-6.6; freeway_mesoscopic_meaning_quick_scan.md, sections 4-5.
""")
start = len(s.shapes)
for y in [1.63, 2.14]:
    rect(s, 1.00, y, 11.25, .34, style.ROAD)
    arrow(s, 11.45, y+.06, color=style.GRAY)
rect(s, 1.00, 2.65, 7.00, .34, style.ROAD)
line(s, 8.0, 2.82, 9.28, 2.31, color=style.GRAY, width=2, end=True)
rect(s, 8.38, 2.71, .50, .24, style.AMBER)
group(s, start, "Conceptual exogenous lane reduction")
for y, value in zip([3.35, 3.80, 4.25, 4.70, 5.15, 5.60, 6.05, 6.55], [
    "Anchor: an exogenous freeway lane-capacity reduction.",
    "Primary choices: speed/advisory control or coordinated ramp metering.",
    "Temporary closure: include post-reopening queue recovery.",
    "Joint control and cooperative merging are expansion candidates.",
    "Test strong lower-resolution models before adding traffic state.",
    "Check derivatives; test action rankings in a richer evaluator.",
    "Macro, richer-model and infeasible outcomes remain admissible.",
    "The control task and minimum resolution remain unselected.",
]):
    body(s, y, value, bold=y == 6.55, color=style.TEAL if y == 6.55 else style.INK)


def validate(saved):
    assert len(saved.slides) == 7
    counts = []
    for index, s in enumerate(saved.slides):
        lines, titles = 0, 0
        assert str(s.background.fill.fore_color.rgb) == style.WHITE
        for shape, (x, y, w, h) in style.leaves(s.shapes):
            assert min(x, y) >= -.01 and x+w <= 13.343 and y+h <= 7.51, shape.name
            if not shape.has_text_frame or not shape.text:
                continue
            title = shape.name.startswith("title:")
            titles += int(title)
            lines += 0 if title else len(shape.text.splitlines())
            for p in shape.text_frame.paragraphs:
                assert p.font.name == style.FONT
                assert p.font.size.pt == (32 if title else 18)
                f = style.font_for(p.font.size.pt, bool(p.font.bold), 144)
                width = (w-(shape.text_frame.margin_left+shape.text_frame.margin_right)/style.EMU)*144
                assert f.getlength(p.text) <= width, (index+1, shape.name, "text width")
            height = len(shape.text.splitlines())*(32 if title else 18)*1.15/72
            available = h-(shape.text_frame.margin_top+shape.text_frame.margin_bottom)/style.EMU
            assert height <= available+.01, (index+1, shape.name, "text height")
        assert titles == 1 and lines <= 8, (index+1, titles, lines)
        assert len(s.notes_slide.notes_text_frame.text) > 200
        counts.append(lines)
    visible = "\n".join(shape.text for s in saved.slides for shape, _ in style.leaves(s.shapes) if shape.has_text_frame)
    assert not re.search(r"\b(RM[0-5]|S0[1-6]|UG1|N1|CQ[1-7]|B1-P|Phase [A-F]|handover|progress)\b", visible, re.I)
    with zipfile.ZipFile(OUT) as z:
        assert z.testzip() is None
    return counts


def contact_sheet(saved):
    # Reuse the style deck's saved-PPTX renderer one page at a time. Its fixed
    # six-page canvas is cropped to the actual page before making this sheet.
    contact = Image.new("RGB", (1350, math.ceil(len(saved.slides)/2)*395+8), "#DCE1E5")
    for i, s in enumerate(saved.slides):
        style.PREVIEW = PREVIEW
        style.preview(SimpleNamespace(slides=[s]))
        with Image.open(PREVIEW) as page:
            thumb = page.crop((8, 8, 675, 383))
        contact.paste(thumb, (8+(i%2)*675, 8+(i//2)*395))
    contact.save(PREVIEW)


if __name__ == "__main__":
    prs.save(OUT)
    saved = Presentation(OUT)
    print("Body-line counts:", validate(saved))
    contact_sheet(saved)
    print(OUT)
    print(PREVIEW)
