"""Build a six-slide, colleague-facing synthesis of the freeway review.

All diagrams are editable grouped PowerPoint shapes. The optional Pillow
preview reads the saved PPTX, but is not an independent Office rendering.
"""

from pathlib import Path
import math
import re

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUT = HERE / "freeway_bottleneck_review_20260918.pptx"
PREVIEW = HERE / "freeway_bottleneck_review_20260918_preview.png"
FONT = "Calibri"
INK = "243341"
TEAL = "267D86"
BLUE = "416F9B"
AMBER = "A87532"
GRAY = "7D8791"
LIGHT = "F1F4F6"
PALE = "EDF5F6"
ROAD = "E4E8EB"
WHITE = "FFFFFF"
EMU = 914400

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Traffic resolution anchored to useful speed advice"
prs.core_properties.subject = "Freeway bottleneck review: methodological findings through 18 September 2026"
prs.core_properties.author = "Research review"


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, x, y, w, h, value, *, title=False, bold=False,
         color=INK, center=False, plot=False, middle=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = ("title:" if title else "plot:" if plot else "body:") + value
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = Inches(.04)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.space_before = p.space_after = Pt(0)
        p.font.name = FONT
        p.font.size = Pt(32 if title else 18)
        p.font.bold = bold or title
        p.font.color.rgb = rgb(color)
    return shape


def rect(slide, x, y, w, h, fill=LIGHT, line=None, *, kind=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.2)
    else:
        s.line.fill.background()
    return s


def node(slide, x, y, w, label, *, color=TEAL, h=.95):
    rect(slide, x, y, w, h, PALE if color == TEAL else LIGHT, color)
    text(slide, x+.05, y+.15, w-.1, h-.3, label, bold=True, center=True, middle=True)


def arrow(slide, x, y, w=.46, *, left=False, color=GRAY):
    return rect(slide, x, y, w, .22, color,
                kind=MSO_SHAPE.LEFT_ARROW if left else MSO_SHAPE.RIGHT_ARROW)


def line(slide, x1, y1, x2, y2, *, color=GRAY, width=1.5, end=False):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if end:
        e = OxmlElement("a:tailEnd")
        e.set("type", "triangle")
        s._element.spPr.get_or_add_ln().append(e)
    return s


def group_since(slide, start, name):
    g = slide.shapes.add_group_shape(list(slide.shapes)[start:])
    g.name = name
    return g


def new_slide(title, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    text(s, .65, .40, 12.0, .68, title, title=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


# 1: motivation, rather than a title/contents boilerplate slide.
s = new_slide("Anchor traffic resolution to a control decision", """Purpose: move from a model-first development sequence to task-anchored methodological design. Starting question: what traffic resolution is necessary for credible, useful control gradients? The current urban simulator is not a fixed constraint on the independent freeway branch. The top sequence presents the risks of choosing detail before evidence; it does not claim earlier implementation was invalid. The bottom sequence states the conditional working line: task and strong baseline, hidden-state decision value, then justified state refinement and separate gradient tests.
Sources: ../differentiable_freeway_bottleneck_task_pre_review/pre_review_closure_and_handoff.md, sections3 P1-P7; ../differentiable_freeway_bottleneck_task_review/01_review_recap_through_s06.md, section1.
""")
text(s, .85, 1.48, 11.7, .40,
     "Model-first risks unnecessary detail and a task chosen to fit it.", bold=True, color=GRAY)
start = len(s.shapes)
for x, label in [(1.0, "Choose model"), (5.0, "Implement"), (9.0, "Find task")]:
    node(s, x, 2.15, 3.3, label, color=GRAY)
arrow(s, 4.45, 2.51); arrow(s, 8.45, 2.51)
group_since(s, start, "Model-first sequence")
text(s, .85, 4.00, 11.7, .40,
     "Working line: establish decision value before choosing richer traffic state.", bold=True, color=TEAL)
start = len(s.shapes)
for x, label in [(1.0, "Task + baseline"), (5.0, "Hidden-state value?"), (9.0, "State + gradients")]:
    node(s, x, 4.75, 3.3, label)
arrow(s, 4.45, 5.11, color=TEAL); arrow(s, 8.45, 5.11, color=TEAL)
group_since(s, start, "Task-to-information-value-to-justified-refinement working line")
s.notes_slide.notes_text_frame.text += "\nWorking-line amendment: the lower sequence replaces the earlier construction diagram. The control task defines the comparison; the strongest aggregate alternative determines what must be hidden; a useful decision difference, not richer forward behavior alone, conditionally justifies added state. Refinement and gradient tests are later obligations, not approved construction. Next slide fixes the task, objective and baseline."

# 2: concrete freeway task and matched lower-resolution alternative.
s = new_slide("A narrowing makes the question concrete", """The operational anchor is an exogenous freeway lane reduction. Initial prospective reference: homogeneous permanent three-to-two narrowing. One infrastructure speed recommendation is shared by all lanes in the advice zone and changes desired free speed; actual motion remains endogenous. No direct vehicle trajectories, forced lane changing or differentiated lane actions. Source mass and waiting must not disappear at the observation boundary. Temporary closure conclusions additionally need reopening and recovery accounting.
Sources: review recap section1; reporting/s01/s01_30_n1_user_gate.md; reporting/s02/s02_35_feedback23_task_and_minimum_information_response.md, section2. Diagram is conceptual, not a simulated road trajectory.
""")
text(s, .85, 1.35, 11.65, .38, "One speed recommendation for all lanes", center=True, bold=True, color=TEAL)
start = len(s.shapes)
for y in [2.18, 3.03]:
    rect(s, 1.0, y, 11.25, .60, ROAD)
    arrow(s, 11.48, y+.19, color=GRAY)
rect(s, 1.0, 3.88, 7.4, .60, ROAD)
line(s, 8.40, 4.18, 9.65, 3.33, end=True, width=2)
rect(s, 3.13, 1.98, 3.40, 2.70, None, TEAL)
line(s, 4.83, 1.80, 4.83, 1.98, color=TEAL, width=2, end=True)
text(s, 8.10, 2.34, 3.85, .36, "Continuing lanes", center=True)
text(s, 1.17, 4.01, 1.92, .36, "Ending lane")
group_since(s, start, "Common-advice road narrowing")
text(s, .85, 5.12, 11.65, .38, "Advice changes desired speed; traffic interactions determine actual motion.")
text(s, .85, 5.61, 11.65, .38, "Strong baseline: spatial traffic, mean response, queues and full history.")
text(s, .85, 6.10, 11.65, .38, "Outcome: total travel and waiting time, including upstream queues.")
text(s, .85, 6.59, 11.65, .38, "Next: what can the strongest aggregate baseline already explain?", bold=True, color=TEAL)
s.notes_slide.notes_text_frame.text += "\nLogical hand-off: common actions and full waiting accounting make the control comparison meaningful. Before attributing value to lane information, review the strongest aggregate response/history alternative; a weak baseline would confound the root question."

# 3: combine framing and close-predecessor reviews into one semantic finding.
s = new_slide("What literature does—and does not—justify", """Combines task framing and predecessor/contrast reviews. Strong aggregate response/history is already credible: Frejo2019 calibrates mean response, Wei2025 retains entry-speed history. Lane transfer models establish plausible behavior, not minimum-information necessity. Seliman2020 macro/micro reoptimization changes objectives, search effort and sometimes actions; Chen2025 lane CTM/MPC compares VSL/no-VSL, not state-only usefulness under the same shared advice. Fixed classes and individual order remain conditional. No first-of-kind or global absence claim.
Sources: reporting/s02/s02_20_n2_findings_and_gate.md sections3-5; reporting/s02/s02_36_feedback23_amendment_novelty_and_component_necessity.md; reporting/s03/s03_20_rb_findings_and_gate.md. Primary anchors: https://doi.org/10.1016/j.trc.2019.01.001 ; https://doi.org/10.1111/mice.13492 ; https://doi.org/10.1061/JTEPBS.0000395 ; https://doi.org/10.1016/j.physa.2025.130876 .
""")
start = len(s.shapes)
rows = [("Aggregate response + history", "Strong baseline—not full obedience"),
        ("Lane distribution + transfer", "Mechanism evidence—not necessity"),
        ("Richer-model control gains", "Check action/objective confounds")]
for i, (a, b) in enumerate(rows):
    y = 1.62 + i*1.22
    node(s, .82, y, 4.75, a, color=BLUE)
    arrow(s, 5.85, y+.37, .52)
    node(s, 6.68, y, 5.83, b)
group_since(s, start, "Predecessor evidence and its limits")
text(s, .85, 5.75, 11.65, .38, "Classes, lanes and individual order are not a mandatory development ladder.")
text(s, .85, 6.30, 11.65, .38, "Next: match full aggregate history, advice choices and objective.", bold=True, color=TEAL)
s.notes_slide.notes_text_frame.text += "\nLogical hand-off: because aggregate response/history already explains important behavior, and richer-control comparisons can change actions or objectives, the next step isolates information value. Equal totals alone are inadequate: the declared retained history must match, with the same advice menu and cost definition. Classes, lanes and order remain conditional rather than assumed necessities."

# 4: grouped illustrative quantitative plot; no fictitious traffic evidence.
s = new_slide("Lane differences need not change decisions", """Illustrative contrast, not a reachable matched-history result. Both totals are30 vehicle-equivalents; ending/continuing group totals are15/15 versus5/25. Continuing lanes are combined for illustration, so this is not a sufficient-state claim. Matching totals alone is weak: the exact-information claim requires equality of the declared strong spatial observation/history channel. Hidden past lane-entry allocations are different unknown past forcing. A particular observer's failure is a distinct model-quality claim. Distinct best actions or a pairwise ranking reversal is insufficient if another common acceptable action exists. Finite-menu evidence does not establish continuous-space or feedback-policy necessity.
Sources: reporting/s04/s04_20_rb_diagnostic_methodology_findings_and_gate.md and recap sections1/S04; current review's lane-contrast explanation. No plot data were measured or simulated.
""")
text(s, .85, 1.30, 11.65, .38, "Illustration only: equal totals, not a certified full-history match.", color=GRAY)
start = len(s.shapes)
for x, label, ending in [(.95, "Situation 1", 15), (7.04, "Situation 2", 5)]:
    text(s, x, 2.00, 5.34, .38, label, center=True, plot=True, bold=True)
    a = 5.34*ending/30
    rect(s, x, 2.65, a, .75, AMBER)
    rect(s, x+a, 2.65, 5.34-a, .75, TEAL)
    text(s, x+.02, 2.78, a-.04, .45, str(ending), plot=True, center=True, bold=True, color=WHITE)
    text(s, x+a+.02, 2.78, 5.30-a, .45, str(30-ending), plot=True, center=True, bold=True, color=WHITE)
rect(s, 3.08, 3.90, .22, .22, AMBER)
text(s, 3.42, 3.85, 2.02, .36, "Ending lane", plot=True)
rect(s, 7.07, 3.90, .22, .22, TEAL)
text(s, 7.40, 3.85, 3.25, .36, "Continuing lanes", plot=True)
group_since(s, start, "Illustrative 30-vehicle-equivalent lane-allocation plot")
text(s, .85, 4.65, 11.65, .38, "Required: match the complete retained aggregate observations and history.", bold=True)
text(s, .85, 5.15, 11.65, .38, "Information loss: identical inputs still conceal decision-relevant differences.")
text(s, .85, 5.65, 11.65, .38, "Model inadequacy: available information is used poorly by one model.")
text(s, .85, 6.15, 11.65, .38, "Decision-loss test: no common acceptable advice in the tested menu.")
text(s, .85, 6.65, 11.65, .38, "Next: can hidden lane state alter the value of common advice?", color=TEAL, bold=True)
s.notes_slide.notes_text_frame.text += "\nLogical hand-off: the matched-history contract states what would count as lost information, but does not establish that such information changes useful advice. The next analytical step must jointly supply hidden contrast and decision relevance, including the possibility of a common acceptable action."

# 5: conditional deductions, not proof of native/reference usefulness.
s = new_slide("Hidden state must change useful decisions", """Conditional original deductions: conservative transfers cancel from aggregate mass in a homogeneous free-flow branch, allowing an early hidden lane contrast without certifying native near-drop preparations. In the separate work-conserving fixed-service subcase, advice-independent service and fastest-arrival dominance make highest advice a common best action, with source/approach waiting counted. This is not universal no-benefit from speed advice. Need endogenous service/receiving/propagation or failure of arrival dominance plus hidden-state dependence and benefit exceeding slowing delay. Native coupling is mechanism evidence only, not a proof that advice is beneficial. A missing held-highest-advice stress probe can leave a common better action outside a pulse menu.
Sources: reporting/s05/s05_20_rb_certificate_findings_and_gate.md; synthesis/s05/rb_certificate_analysis.pdf; reporting/s06/s06_20_reference_bridge_findings_and_gate.md. Native model: Jakob Erdmann2015, https://elib.dlr.de/102254/1/Springer-SUMOs_Lane_changing_model.pdf . No borrowed discrete resets, random mass shifts, smoothing or invented penalty is approved.
""")
start = len(s.shapes)
for x, lab in [(.82, "Shared advice"), (3.90, "Desired speed"), (6.98, "Merging"), (10.06, "Outcome?")]:
    node(s, x, 1.75, 2.45, lab)
for x in [3.40, 6.48, 9.56]:
    arrow(s, x, 2.11, .36, color=TEAL)
group_since(s, start, "Advice-response-interaction-usefulness chain")
text(s, .85, 3.45, 11.65, .46, "Hidden contrast is possible in conditional free flow—but that is not enough.")
text(s, .85, 4.22, 11.65, .46, "Fixed-service subcase: earlier arrivals make highest advice a common best.")
text(s, .85, 4.99, 11.65, .46, "Thus inspect native merging: can state-dependent benefits outweigh delay?")
text(s, .85, 5.76, 11.65, .46, "Next: find a reachable hidden-state contrast with different useful advice.", bold=True, color=TEAL)
s.notes_slide.notes_text_frame.text += "\nLogical hand-off: the hidden-state construction and fixed-service obstruction are separate conditional deductions, not a combined native certificate. They explain why the subsequent reference-first review inspects endogenous merge interactions rather than inserting an arbitrary penalty. A credible mechanism must still be instantiated in reachable preparations with matched declared history and no common acceptable advice."

# 6: open methodological boundary, without user gates or project workflow.
s = new_slide("Mechanisms exist; information value is open", """Current synthesis: no complete positive or negative conclusion for the joint information-to-useful-advice claim. Advisory aggregate-feedback benefits (FB0005) and occupancy-conditioned lane fractions (FB0039) strengthen the aggregate alternative, without proving sufficiency. Native gap/speed/cooperation and field redistribution are precedents, not minimum-state or useful-gradient evidence. Strong regulatory low-speed lane-drop outcomes (Zhang2023) do not transfer automatically to moderate voluntary advice. Preparation must be reachable and survive native anticipation; later interaction may expose the difference in mean-speed/history. The proposed40/60s times cannot reach the final400m under intended road-start/speed bounds;80s remains sparse under conditional entry defaults. No server/simulation result was used. Diagram shows the semantic tension between earlier hidden information and near-merge consequences, not an approved scenario change.
Sources: reporting/s06/s06_20_reference_bridge_findings_and_gate.md; extraction/s06_reference_mechanism_source_audit.md; recap section3. New primary lane-drop manuscript: https://ira.lib.polyu.edu.hk/bitstream/10397/92568/1/Zhang_Optimization_Speed_Freeway.pdf . Future minimal-refinement, forward, internal gradient and independent usefulness obligations remain separate. Review/communication does not authorize construction or experiments.
""")
text(s, .85, 1.35, 11.65, .40, "Supported: strong aggregate alternatives and plausible merge mechanisms.")
text(s, .85, 1.88, 11.65, .40, "Missing: a reachable full-history match requiring different useful advice.", bold=True, color=TEAL)
start = len(s.shapes)
for y in [2.80, 3.28]:
    rect(s, 1.00, y, 11.3, .29, ROAD)
rect(s, 1.00, 3.76, 7.96, .29, ROAD)
line(s, 8.96, 3.90, 10.10, 3.42, end=True)
for x in [3.40, 8.30]:
    rect(s, x-.48, 2.63, .96, 1.62, None, TEAL, kind=MSO_SHAPE.OVAL)
text(s, 2.04, 4.40, 2.72, .40, "Earlier upstream", center=True)
text(s, 6.67, 4.40, 3.26, .40, "Near the narrowing", center=True)
group_since(s, start, "Where hidden lane contrast can survive")
text(s, .85, 5.17, 11.65, .40, "Next: make the contrast reachable without omitting informative history.", bold=True, color=TEAL)
text(s, .85, 5.72, 11.65, .40, "If useful advice must differ, test the smallest conservative added state.")
text(s, .85, 6.27, 11.65, .40, "Then test gradient correctness and independent usefulness separately.")
text(s, .85, 6.82, 11.65, .40, "Without that evidence, do not infer that richer traffic detail is necessary.", bold=True, color=TEAL)
s.notes_slide.notes_text_frame.text += "\nLogical synthesis: source-supported mechanisms do not close the joint witness. The remaining review step concerns reachable contrast preparation under the full retained-history contract. Only a demonstrated decision-value witness can motivate the smallest conservative refinement; forward fidelity, gradient correctness and independent control usefulness then require distinct evidence. Failure to find a witness is not a global sufficiency theorem. Current review gate remains open; this deck amendment approves neither a preparation change nor any experiment."


def leaves(shapes, ox=0, oy=0, sx=1, sy=1):
    """Flatten groups with their actual OOXML child-coordinate transforms."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xf = shape._element.grpSpPr.xfrm
            co, ce = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
            gx, gy = shape.width/int(ce.get("cx")), shape.height/int(ce.get("cy"))
            yield from leaves(shape.shapes,
                              ox+sx*(shape.left-int(co.get("x"))*gx),
                              oy+sy*(shape.top-int(co.get("y"))*gy), sx*gx, sy*gy)
        else:
            yield shape, ((ox+sx*shape.left)/EMU, (oy+sy*shape.top)/EMU,
                          sx*shape.width/EMU, sy*shape.height/EMU)


def font_for(points, bold, scale=100):
    path = "/usr/share/fonts/truetype/dejavu/DejaVuSans" + ("-Bold" if bold else "") + ".ttf"
    return ImageFont.truetype(path, round(points*scale/72))


def color_of(color, default):
    try:
        return "#" + str(color.rgb)
    except (AttributeError, ValueError, TypeError):
        return default


def check_deck(saved):
    assert len(saved.slides) == 6
    counts = []
    for slide in saved.slides:
        body_lines, titles = 0, 0
        for shape, (x, y, w, h) in leaves(slide.shapes):
            assert min(x, y) >= -.01 and x+w <= 13.343 and y+h <= 7.51, shape.name
            if not shape.has_text_frame or not shape.text:
                continue
            is_title = shape.name.startswith("title:")
            titles += int(is_title)
            if not is_title and not shape.name.startswith("plot:"):
                body_lines += len(shape.text.splitlines())
            for p in shape.text_frame.paragraphs:
                assert p.font.name == FONT
                assert p.font.size.pt == (32 if is_title else 18)
                font = font_for(p.font.size.pt, bool(p.font.bold), 144)
                available = (w-(shape.text_frame.margin_left+shape.text_frame.margin_right)/EMU)*144
                assert font.getlength(p.text) <= available, (shape.name, "width overflow")
            required = len(shape.text.splitlines())*(32 if is_title else 18)*1.15/72
            available_h = h-(shape.text_frame.margin_top+shape.text_frame.margin_bottom)/EMU
            assert required <= available_h+.01, (shape.name, "height overflow")
        assert titles == 1 and body_lines <= 8, (titles, body_lines)
        assert len(slide.notes_slide.notes_text_frame.text) > 200
        counts.append(body_lines)
    visible = "\n".join(shape.text for slide in saved.slides
                        for shape, _ in leaves(slide.shapes) if shape.has_text_frame)
    assert not re.search(r"\b(S0[1-6]|DG-|UG1|Candidate [ABC]|Phase [A-F])", visible)
    assert not any(slide.shapes.title for slide in saved.slides)
    return counts


def preview(saved):
    """A font-substituted layout proof from the saved PPTX, not Office fidelity."""
    scale = 100
    pages = []
    for slide in saved.slides:
        im = Image.new("RGB", (1334, 750), "#"+WHITE)
        draw = ImageDraw.Draw(im)
        for shape, (x, y, w, h) in leaves(slide.shapes):
            x, y, w, h = (v*scale for v in (x, y, w, h))
            bounds = (x, y, x+w, y+h)
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                # a connector can flip vertically/horizontally in its transform.
                xf = shape._element.spPr.xfrm
                xx = (x+w, x) if xf.get("flipH") == "1" else (x, x+w)
                yy = (y+h, y) if xf.get("flipV") == "1" else (y, y+h)
                col = color_of(shape.line.color, "#"+GRAY)
                draw.line((xx[0], yy[0], xx[1], yy[1]), fill=col, width=3)
                if shape._element.spPr.ln.find(qn("a:tailEnd")) is not None:
                    ang = math.atan2(yy[1]-yy[0], xx[1]-xx[0])
                    tip = (xx[1], yy[1])
                    draw.polygon([tip, (tip[0]-12*math.cos(ang-.35), tip[1]-12*math.sin(ang-.35)),
                                  (tip[0]-12*math.cos(ang+.35), tip[1]-12*math.sin(ang+.35))], fill=col)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = color_of(shape.fill.fore_color, None) if shape.fill.type == 1 else None
                try:
                    col = color_of(shape.line.color, None)
                except AttributeError:
                    col = None
                kind = shape.auto_shape_type
                if kind == MSO_SHAPE.OVAL:
                    draw.ellipse(bounds, fill=fill, outline=col, width=2)
                elif kind in [MSO_SHAPE.LEFT_ARROW, MSO_SHAPE.RIGHT_ARROW]:
                    points = [(x,y+h*.25), (x+w*.62,y+h*.25), (x+w*.62,y),
                              (x+w,y+h*.5), (x+w*.62,y+h), (x+w*.62,y+h*.75), (x,y+h*.75)]
                    if kind == MSO_SHAPE.LEFT_ARROW:
                        points = [(2*x+w-px, py) for px, py in points]
                    draw.polygon(points, fill=fill, outline=col)
                else:
                    draw.rectangle(bounds, fill=fill, outline=col, width=2)
            if not shape.has_text_frame or not shape.text:
                continue
            frame = shape.text_frame
            ml, mt = frame.margin_left/EMU*scale, frame.margin_top/EMU*scale
            mr, mb = frame.margin_right/EMU*scale, frame.margin_bottom/EMU*scale
            ps = list(frame.paragraphs)
            pts = ps[0].font.size.pt
            step = pts*scale/72*1.15
            ty = y+mt
            if frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
                ty = y+mt+(h-mt-mb-len(ps)*step)/2
            for p in ps:
                font = font_for(p.font.size.pt, bool(p.font.bold), scale)
                tx = x+ml
                if p.alignment == PP_ALIGN.CENTER:
                    tx += (w-ml-mr-font.getlength(p.text))/2
                draw.text((tx, ty), p.text, font=font, fill=color_of(p.font.color, "#"+INK), anchor="lt")
                ty += step
        pages.append(im)
    contact = Image.new("RGB", (1350, 1185), "#DCE1E5")
    for i, im in enumerate(pages):
        thumb = im.resize((667, 375), Image.Resampling.LANCZOS)
        contact.paste(thumb, (8+(i%2)*675, 8+(i//2)*395))
    contact.save(PREVIEW)


if __name__ == "__main__":
    prs.save(OUT)
    saved = Presentation(OUT)
    print("Body-line counts, excluding the illustrative plot:", check_deck(saved))
    preview(saved)
    print(OUT)
    print(PREVIEW)
