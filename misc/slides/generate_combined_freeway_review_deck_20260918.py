"""Merge the saved historical and continuation decks, preserving native shapes.

No source presentation is rebuilt or edited. Media/hyperlink relationships are
explicitly rejected because these two decks contain self-contained native shapes.
"""

from copy import deepcopy
from hashlib import sha256
from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from types import SimpleNamespace
import math
import re
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
EARLY = HERE / "freeway_review_to_task_first_handoff_20260918.pptx"
LATER = HERE / "freeway_bottleneck_review_20260918.pptx"
OUT = HERE / "combined_freeway_review_20260918.pptx"
PREVIEW = HERE / "combined_freeway_review_20260918_preview.png"
spec = spec_from_file_location("freeway_style", HERE / "generate_freeway_review_deck_20260918.py")
style = module_from_spec(spec)
spec.loader.exec_module(style)
style.prs = Presentation()
prs = style.prs
prs.slide_width = style.Inches(13.333333)
prs.slide_height = style.Inches(7.5)
prs.core_properties.title = "Traffic resolution for useful freeway control"
prs.core_properties.subject = "From model compatibility to task-first information-value findings through 18 September 2026"
prs.core_properties.author = "Research review"
early, later = Presentation(EARLY), Presentation(LATER)
assert len(early.slides) == 7 and len(later.slides) == 6
assert (early.slide_width, early.slide_height) == (later.slide_width, later.slide_height)
SOURCE_HASHES = {p.name: sha256(p.read_bytes()).hexdigest() for p in (EARLY, LATER)}
unchanged_pairs = []


def clone(source):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = source._element.cSld.find(qn("p:bg"))
    assert bg is not None
    s._element.cSld.insert(0, deepcopy(bg))
    for shape in source.shapes:
        for element in shape._element.iter():
            assert not any(k.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") for k in element.attrib), "Relationship-aware copying required"
        s.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    s.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text
    unchanged_pairs.append((source, s))
    return s


def replace(s, old, new):
    matches = [sh for sh, _ in style.leaves(s.shapes) if sh.has_text_frame and sh.text == old]
    assert len(matches) == 1, old
    sh = matches[0]
    assert len(sh.text_frame.paragraphs) == 1
    sh.text_frame.paragraphs[0].text = new
    prefix = sh.name.split(":", 1)[0]
    sh.name = prefix + ":" + new


s = style.new_slide("Traffic resolution for useful freeway control", """Combined title slide. This presentation starts with the original compatibility-framed ramp-metering feasibility review and continues through the later independent task-first freeway review and its information/decision-value findings. The scientific endpoint is the continuation deck's 18 September 2026 synthesis, not the earlier pre-review handover. The original bounded conclusion is retained, while later conclusions remain conditional and open where the source deck says so. This merge changes communication only: no new experiment, implementation, scientific claim or review-gate decision is introduced.
Inputs: freeway_review_to_task_first_handoff_20260918.pptx and freeway_bottleneck_review_20260918.pptx. See companion notes for exact slide mapping and source checksums.
""")
style.text(s, .85, 2.42, 11.65, .48, "From model compatibility to task-first information value", bold=True, color=style.TEAL)
style.text(s, .85, 3.25, 11.65, .42, "Ramp-control evidence, common speed advice and the necessary-resolution question")
style.text(s, .85, 4.10, 11.65, .42, "Methodological findings through 18 September 2026")
style.text(s, .85, 5.90, 11.65, .42, "Technical discussion for colleagues")

s = style.new_slide("Contents", """Combined contents. Two boilerplate and nine substantive slides. Slides 3-5 preserve the historical evidence and restricted transfer-case findings. Slide 6 merges the historical scope correction with the continuation's task-first working line; slide 7 supplies the narrowed speed-advice task in place of the earlier provisional candidate inventory. Slides 8-11 continue with close-predecessor evidence, the matched-history decision test, conditional analytical deductions and the unresolved reachable information-value contrast. All visible content concerns methodology and findings. Historical pre-review choices are not presented as the current unresolved task selection; later conditional deductions are not presented as empirical traffic results or universal sufficiency theorems.
Detailed original sources and their limitations remain in the preserved slide notes and source presentations' companion records.
""")
for y, value in zip([1.70, 2.60, 3.50, 4.40, 5.30], [
    "3–5     Strong alternatives and the bounded ordered-stream result",
    "6–7     Task-first framing and common speed advice at a narrowing",
    "8          What the closest control literature supports",
    "9          Separate information loss from model inadequacy",
    "10–11  Hidden-state decision value: findings and remaining gap",
]):
    style.text(s, .85, y, 11.65, .42, value)

for source in list(early.slides)[2:5]:
    clone(source)

# Replace the overlapping old slides 6/7 with the later working line and task.
s = clone(later.slides[0])
unchanged_pairs.pop()
replace(s, "Anchor traffic resolution to a control decision", "Task validity precedes model compatibility")
replace(s, "Model-first risks unnecessary detail and a task chosen to fit it.", "The earlier negative result was scoped to the fixed architecture.")
replace(s, "Choose model", "Fixed model")
replace(s, "Implement", "Compatible task")
replace(s, "Find task", "Value within scope")
replace(s, "Working line: establish decision value before choosing richer traffic state.", "Task-first: establish decision value before choosing richer traffic state.")
s.notes_slide.notes_text_frame.text += """
Merged transition amendment: the gray sequence now names the actual compatibility-framed question in the preceding historical slides, rather than repeating a generic choose-model/implement/find-task sequence. Its bounded negative result remains valid within that scope. The teal sequence establishes the independent task-first evaluation: strongest task baseline, decision value of hidden state, then conditional state refinement and separate gradient checks. The urban and freeway branches share the resolution question and evidence logic; they need not share a model, codebase or paper. Compatibility and reuse remain later secondary audits. This replaces the old historical slides 6 and 7's transition/provisional-task endpoint with a coherent lead-in to the subsequent narrowed common-speed-advice task.
Additional provenance: ../differentiable_freeway_bottleneck_task_pre_review/feedback22_email_trace_response.md, opening clarification and sections 6.1/6.7; pre_review_closure_and_handoff.md, P1-P8; ../../temp_content/feedback/email_trace_20260916.md, messages 2 and 4.
"""

for source in list(later.slides)[1:]:
    clone(source)


def validate(saved):
    assert len(saved.slides) == 11
    counts = []
    for i, s in enumerate(saved.slides):
        titles, lines = 0, 0
        assert str(s.background.fill.fore_color.rgb) == style.WHITE
        for sh, (x, y, w, h) in style.leaves(s.shapes):
            assert min(x, y) >= -.01 and x+w <= 13.343 and y+h <= 7.51, sh.name
            if not sh.has_text_frame or not sh.text:
                continue
            title = sh.name.startswith("title:")
            titles += int(title)
            if not title and not sh.name.startswith("plot:"):
                lines += len(sh.text.splitlines())
            for p in sh.text_frame.paragraphs:
                assert p.font.name == style.FONT
                assert p.font.size.pt == (32 if title else 18)
                f = style.font_for(p.font.size.pt, bool(p.font.bold), 144)
                available = (w-(sh.text_frame.margin_left+sh.text_frame.margin_right)/style.EMU)*144
                assert f.getlength(p.text) <= available, (i+1, sh.name, "width")
            required = len(sh.text.splitlines())*(32 if title else 18)*1.15/72
            available_h = h-(sh.text_frame.margin_top+sh.text_frame.margin_bottom)/style.EMU
            assert required <= available_h+.01, (i+1, sh.name, "height")
        assert titles == 1 and lines <= 8, (i+1, titles, lines)
        assert len(s.notes_slide.notes_text_frame.text) > 200
        counts.append(lines)
    # Eight unchanged content slides retain every native shape and original note.
    assert len(unchanged_pairs) == 8
    for source, target in unchanged_pairs:
        idx = list(prs.slides).index(target)
        actual = saved.slides[idx]
        assert [sh._element.xml for sh in actual.shapes] == [sh._element.xml for sh in source.shapes]
        assert actual.notes_slide.notes_text_frame.text == source.notes_slide.notes_text_frame.text
    visible = "\n".join(sh.text for s in saved.slides for sh, _ in style.leaves(s.shapes) if sh.has_text_frame)
    assert "remain unselected" not in visible
    assert not re.search(r"\b(RM[0-5]|S0[1-6]|UG1|N1|CQ[1-7]|B1-P|Phase [A-F]|handover|progress)\b", visible, re.I)
    with zipfile.ZipFile(OUT) as archive:
        assert archive.testzip() is None
    assert SOURCE_HASHES == {p.name: sha256(p.read_bytes()).hexdigest() for p in (EARLY, LATER)}
    return counts


def preview(saved):
    contact = Image.new("RGB", (1350, math.ceil(len(saved.slides)/2)*395+8), "#DCE1E5")
    for i, s in enumerate(saved.slides):
        style.PREVIEW = PREVIEW
        style.preview(SimpleNamespace(slides=[s]))
        with Image.open(PREVIEW) as page:
            thumb = page.crop((8, 8, 675, 383))
        contact.paste(thumb, (8+(i%2)*675, 8+(i//2)*395))
    contact.save(PREVIEW)


if __name__ == "__main__":
    prs.save(OUT)
    saved = Presentation(OUT)
    print("Non-title lines excluding plot labels:", validate(saved))
    preview(saved)
    print("Source SHA-256:", SOURCE_HASHES)
    print(OUT)
    print(PREVIEW)
