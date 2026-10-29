"""Request 7: weekly task-definition review, through 9 October 2026.

Fourteen slides: two boilerplate and twelve substantive slides.
Native grouped diagrams; typeset SVG mathematics with embedded PNG fallbacks.
The layout proof uses DejaVu Sans, not an Office/Calibri rendering.
"""
from pathlib import Path
from io import BytesIO
import json
import math
import os
import posixpath
import re
import zipfile

os.environ.setdefault("MPLCONFIGDIR", "/tmp/freeway_slides_matplotlib")
os.environ.setdefault("XDG_CACHE_HOME", "/tmp/freeway_slides_font_cache")
Path(os.environ["XDG_CACHE_HOME"]).mkdir(parents=True, exist_ok=True)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
OUT = HERE / "weekly_freeway_control_review_20261009.pptx"
ASSETS = HERE / "weekly_freeway_control_review_20261009_assets"
NOTES = HERE / "weekly_freeway_control_review_20261009_notes.md"
REPORT = HERE / "weekly_freeway_control_review_20261009_validation.json"
PREVIEW = HERE / "weekly_freeway_control_review_20261009_preview.png"
ASSETS.mkdir(exist_ok=True)
FONT = "Calibri"
INK, TEAL, BLUE, AMBER = "243341", "267D86", "416F9B", "A87532"
GRAY, LIGHT, ROAD, WHITE = "7D8791", "F1F4F6", "E4E8EB", "FFFFFF"
EMU = 914400
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
VSL = "misc/differentiable_freeway_bottleneck_task_review"
RAMP = "misc/selective_discreteness_ramp_metering_review"
svg_sequence = 0
records = []
math_lines = {}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = "Ramp metering: from hidden information to an adequate control model"
prs.core_properties.subject = "Method-first ramp-metering task and model-reduction review, through 9 October 2026"
prs.core_properties.author = "Research review"
prs.core_properties.keywords = "VSL, ramp metering, hidden information, useful gradients"
plt.rcParams.update({"svg.fonttype": "path", "mathtext.fontset": "cm",
                     "font.family": "DejaVu Sans"})


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, x, y, w, h, value, *, role="body", bold=False,
         color=INK, center=False, middle=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = role + ":" + value.replace("\n", " / ")
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(.035)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if middle else MSO_ANCHOR.TOP
    for i, line in enumerate(value.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.font.name = FONT
        p.font.size = Pt(32 if role == "title" else 18)
        p.font.bold = bold or role == "title"
        p.font.color.rgb = rgb(color)
        p.space_before = p.space_after = Pt(0)
    return shape


def rect(slide, x, y, w, h, fill=LIGHT, line=None, kind=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.3)
    else:
        s.line.fill.background()
    return s


def arrow(slide, x, y, w=.5, h=.22, color=GRAY, kind=MSO_SHAPE.RIGHT_ARROW):
    return rect(slide, x, y, w, h, color, kind=kind)


def line(slide, x1, y1, x2, y2, color=GRAY, width=1.7, end=False):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if end:
        e = OxmlElement("a:tailEnd")
        e.set("type", "triangle")
        s._element.spPr.get_or_add_ln().append(e)
    return s


def group(slide, start, name):
    s = slide.shapes.add_group_shape(list(slide.shapes)[start:])
    s.name = "diagram:" + name
    return s


def slide(title, explanation, sources, *, boilerplate=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    text(s, .65, .4, 12.02, .72, title, role="title")
    source_text = "\n".join(sources)
    s.notes_slide.notes_text_frame.text = explanation + "\n\nSources:\n" + source_text
    records.append({"slide": len(prs.slides), "title": title,
                    "boilerplate": boilerplate, "notes": explanation,
                    "sources": sources})
    return s


def body(slide, values, y=5.1, step=.48, x=.82, w=11.75):
    for i, value in enumerate(values):
        text(slide, x, y + step*i, w, .40, value)


def vector_picture(slide, name, svg, png, x, y, w, h):
    """Verified Part/relationship API; SVG blip plus normal PNG image part."""
    global svg_sequence
    with Image.open(png) as image:
        iw, ih = image.size
    factor = min(w/iw, h/ih)
    ww, hh = iw*factor, ih*factor
    pic = slide.shapes.add_picture(str(png), Inches(x+(w-ww)/2),
                                  Inches(y+(h-hh)/2), Inches(ww), Inches(hh))
    pic.name = name
    svg_sequence += 1
    part = Part(PackURI(f"/ppt/media/weekly_vector_{svg_sequence}.svg"),
                "image/svg+xml", slide.part.package, Path(svg).read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.find(qn("p:blipFill")).find(qn("a:blip"))
    extlist = blip.find(qn("a:extLst"))
    if extlist is None:
        extlist = OxmlElement("a:extLst")
        blip.append(extlist)
    ext = OxmlElement("a:ext")
    ext.set("uri", SVG_EXT)
    node = etree.Element(f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS})
    node.set(qn("r:embed"), rid)
    ext.append(node)
    extlist.append(ext)
    return pic


def formula(slide, name, lines, x, y, w, h, fontsize=25, line_spacing=.56):
    fig = plt.figure(figsize=(11.6, line_spacing*len(lines)+.18), dpi=240)
    fig.patch.set_alpha(0)
    for i, value in enumerate(lines):
        fig.text(.5, 1-(i+.6)/len(lines), "$"+value+"$",
                 ha="center", va="center", fontsize=fontsize, color="#"+INK)
    svg, png = ASSETS / (name+".svg"), ASSETS / (name+".png")
    fig.savefig(svg, transparent=True, bbox_inches="tight", pad_inches=.045)
    fig.savefig(png, transparent=True, bbox_inches="tight", pad_inches=.045)
    plt.close(fig)
    shape = vector_picture(slide, "math:"+name, svg, png, x, y, w, h)
    math_lines[shape.name] = len(lines)


def chain(slide, labels, *, y=1.62):
    start = len(slide.shapes)
    span, gap = 11.8, .35
    width = (span-gap*(len(labels)-1))/len(labels)
    for i, label in enumerate(labels):
        x = .78 + i*(width+gap)
        text(slide, x, y, width, .77, label, role="plot",
             bold=True, color=TEAL, center=True, middle=True)
        if i < len(labels)-1:
            arrow(slide, x+width+.025, y+.30, gap-.05, color=TEAL)
    group(slide, start, "causal sensitivity chain")


def lead(s, value):
    text(s, .80, 1.38, 11.85, .48, value, bold=True)


def label(s, x, y, w, value, color=INK, bold=False, h=.42, center=False):
    return text(s, x, y, w, h, value, role="plot", color=color,
                bold=bold, center=center)


def density_strip(s, x, y, w, occupied, where="rear", color=BLUE):
    """A density region, not a sequence of individual vehicles."""
    rect(s, x, y, w, .47, WHITE, GRAY)
    xx = x if where == "rear" else x+w-occupied
    rect(s, xx, y+.035, occupied, .40, color)
    arrow(s, x+w+.12, y+.14, .36, .18, TEAL)


R = RAMP
F = R+"/reporting/rm12/bp_r_stage3/"
SR = R+"/synthesis/rm12/bp_r_stage3/"
EX = R+"/extraction/rm12/bp_r_stage3/"

# Slides 1--2: exactly two boilerplate pages.
s = slide("Ramp metering: from detail to decisions",
    "Request 7, week ending 9 October 2026. Continue from last week's mass/flow "
    "information-hiding obstruction through the completed D2 forward-contract "
    "continuation. Slides 3--10 are the group-meeting backbone; 11--14 are "
    "technical discussion pages. This deck synthesizes existing source/static "
    "review, not a simulation or a new scientific pass. The review is temporarily "
    "parked at Q.J.F.R; preserving the full field is recommended but unapproved.",
    ["temp_content/feedback/feedback.md", F+"d2_forward/findings_and_field_gate.md"],
    boilerplate=True)
text(s, .85, 2.15, 11.65, .62, "Which retained traffic state can support useful rate decisions?", bold=True)
text(s, .85, 3.17, 11.65, .55, "Weekly methodology review · 9 October 2026", color=TEAL)
text(s, .85, 4.17, 11.65, .95,
     "One local ramp merge · common rate decisions · complete system time\n"
     "Source review and analytical checks; the research task remains open.")

s = slide("Contents",
    "The primary presentation is eight substantive slides, 3--10. It gives a "
    "complete semantic review arc: the original obstruction, bounded negative "
    "mechanism screens, operational-information backtracking, revised "
    "representation claim, supported physical family, memory/exact-state "
    "construction and its remaining gaps. Four further substantive slides "
    "supply technical examples for expert conversation without adding slides "
    "beyond the requested twelve-page substantive limit.",
    ["temp_content/feedback/feedback.md"], boilerplate=True)
body(s, ["Group meeting · slides 3–10", "Hidden information: why the original chain did not close",
         "Model adequacy: what the changed claim now lets us construct",
         "Technical discussion · slides 11–14",
         "Creation and recovery; shared service; spatial fields; control evidence"],
     y=2.0, step=.72)

# Eight primary pages: a stand-alone group-meeting story.
s = slide("Which traffic detail matters for control?",
    "Begin with the question left open last week, not a progress log. Originally "
    "we sought traffic detail unresolved after all admitted aggregate mass, "
    "flow, supported speed, ramp and control history, which also changed useful "
    "common ramp-rate actions under complete all-stream TST. This week checked "
    "why that task did not follow from physical merge effects, then explicitly "
    "backtracked to model-reduction adequacy. The general research motivation "
    "survives, but the original hidden-history claim is not established. Forward "
    "fidelity, derivative correctness and useful decisions are separate. A "
    "compact sufficient model is an acceptable outcome. Use the source and "
    "small analytical certificates before any implementation or experiment.",
    [R+"/reporting/rm9_99_scoped_closure.md",
     R+"/reporting/rm12/bp_r/semantic_position_in_review_tree.md",
     F+"baseline_review_scope_check.md"])
lead(s, "The goal is a justified chain from traffic detail to useful rate decisions.")
body(s, ["1. What traffic detail survives the full declared aggregate history?",
         "2. Does retaining that detail change useful common ramp-rate decisions?",
         "3. Can its gradient predict lower complete system time?"],
     y=2.20, step=.80)
body(s, ["Keep genuine tasks and strong rivals; add assumptions only with justification.",
         "Separate physical support, state adequacy, control relevance and useful gradients."],
     y=5.45, step=.57)

s = slide("Last week made full history the obstacle",
    "Recap the last deck's obstruction, then this week's stronger-history audit. "
    "The normalized two-lane free-branch example has known coefficients c1=1, "
    "c2=1.5. Masses (2,8) and (6,4) both total 10 but yield flows 14 and 12; "
    "in that branch total mass and aggregate flow reconstruct the split. These "
    "are normalized algebraic units, not calibrated traffic observations. "
    "Subsequent RM9 observer-C work uses section/Edie-supported speed accounting "
    "rather than inventing a favorable speed channel. If a complete earlier "
    "physical state is uniquely known and the dynamics, controls and forcing "
    "are exact and public, deterministic evolution fixes the future. Knowing "
    "only aggregate initial statistics is different; congested/history domains "
    "were not universally ruled out. Neither equal snapshots nor extra lanes "
    "alone certify a lawful full-history/action distinction.",
    ["misc/slides/weekly_freeway_control_review_20261002_notes.md",
     R+"/reporting/rm9_99_scoped_closure.md",
     R+"/synthesis/rm9/rm9_separate_resource_and_observer_audit.tex"])
lead(s, "Equal density alone was too weak: flow and earlier knowledge can reveal the split.")
chain(s,["Last week's obstacle\nMass + flow reveal split",
         "Required evidence\nDetail survives history",
         "This week's search\nGenuine mechanisms"],y=2.65)
body(s,["A known complete past state and exact public dynamics can determine later detail.",
        "We therefore need more than different lane snapshots to establish useful hiddenness.",
        "This is a conditional obstruction; aggregate initial statistics need not reveal everything."],
     y=4.56,step=.62)

s = slide("Real mechanisms do not prove hidden detail",
    "RM10 reviewed acceleration/recovery: recovery means how a slow entrant "
    "accelerates, encounters traffic and ceases to impose an extra restriction, "
    "not that the whole freeway is uncongested. Sources support the mechanism. "
    "Some mean/speed/hysteresis models are conditionally sufficient; richer "
    "families did not bind genuine initialization, equal full history and "
    "different useful common rates. RM11 C0 separately reviewed finite "
    "longitudinal merging opportunity with genuine sites, including a two-lane "
    "antecedent. Public lane geometry and observed position can be retained by "
    "spatial macroscopic competitors. Residual gap/position information and "
    "its action relevance remained unestablished. Both bounded reviews closed "
    "with scoped negative task admission; neither proved the mechanisms false "
    "or all aggregate representations sufficient. The diagrams explain "
    "physical leads, not reached matched-history cases.",
    [R+"/reporting/rm10/rm10_70_findings_and_conclusion.md",
     R+"/reporting/rm11/c0_findings_and_disposition.md"])
lead(s,"The candidate detail was recovery state and merge opportunity, not total mass.")
start=len(s.shapes)
label(s,.92,2.10,5.6,"Recovery detail",bold=True)
label(s,.92,2.65,5.6,"Slow entrants' location + speed",color=AMBER)
density_strip(s,1.05,3.25,4.5,1.55,"rear")
line(s,2.80,3.01,2.80,4.00,AMBER,2.3)
line(s,2.98,3.47,4.12,3.47,AMBER,2.0,True)
label(s,1.05,4.15,5.5,"Remaining acceleration changes service")
label(s,6.85,2.10,5.6,"Merge-opportunity detail",bold=True)
label(s,6.85,2.65,5.6,"Distance to lane end + gap timing",color=TEAL)
rect(s,6.97,3.25,4.65,.36,ROAD)
rect(s,8.80,3.26,.92,.34,WHITE)
rect(s,6.97,3.99,3.10,.36,ROAD)
line(s,9.20,4.15,9.20,3.61,TEAL,2.0,True)
line(s,10.07,3.99,10.07,4.35,GRAY,2.0)
label(s,7.02,4.65,5.2,"Spatial models can retain opportunity")
group(s,start,"candidate recovery state and finite merge opportunity")
body(s,["The sources support these effects on traffic, not their absence from all available history.",
        "Known motion or a spatial aggregate state can already retain part of this information.",
        "Neither reviewed lead supplied the full-history distinction and its useful common-rate effect."],
     y=5.38,step=.48)

s = slide("Backtrack from hiddenness to model adequacy",
    "RM12 Stage 1 BP-I audited a documented operational information domain. "
    "Point-detector inputs are not automatically a complete cell-wide observer, "
    "but used inputs are better documented than all accessible initial/history "
    "knowledge. No justified missing-history witness was admitted. BP-R then "
    "explicitly changed the claim: what response does an independently evolving "
    "model lose when it compresses a shared known scenario? One local ramp, "
    "common bounded upstream rate schedules, public forcing and complete "
    "all-stream TST remain. The slow group's condition may be known to every "
    "model, each initialized by its declared projection. It is then a question "
    "of representation, not private controller information. No repeated feed "
    "of future reference trajectories is allowed. Spatial/memory/moment models "
    "are strong rivals; a sufficient small summary is a positive result. "
    "Backtracking preserves the earlier scoped negatives instead of relabeling "
    "them as a newly solved original information question.",
    [R+"/reporting/rm12/bp_i/findings_and_gate.md",
     R+"/reporting/rm12/bp_r/semantic_position_in_review_tree.md"])
lead(s,"Known traffic can be public yet omitted from a model's retained state.")
start=len(s.shapes)
label(s,.95,2.15,5.5,"First: ground operational knowledge",bold=True)
rect(s,1.1,3.1,4.6,.18,ROAD)
line(s,2.55,2.8,2.55,3.51,TEAL,2.7)
label(s,1.05,3.78,5.35,"Detector inputs ≠ all accessible history")
label(s,1.05,4.33,5.30,"Knowledge contract stayed incomplete",color=AMBER)
line(s,5.90,3.45,6.57,3.45,TEAL,2.1,True)
label(s,6.82,2.15,5.5,"Then: test representation adequacy",bold=True)
density_strip(s,7.0,2.92,4.18,1.5,"rear")
line(s,8.7,2.77,8.7,3.55,AMBER,2)
line(s,8.65,3.60,7.90,4.10,GRAY,1.6,True)
line(s,8.65,3.60,10.48,4.10,GRAY,1.6,True)
rect(s,7.05,4.30,1.92,.33,WHITE,GRAY)
rect(s,7.07,4.32,.72,.29,BLUE)
line(s,7.82,4.12,7.82,4.82,AMBER,2.0)
label(s,6.72,5.09,2.9,"Field + exact effects")
rect(s,10.02,4.30,1.92,.33,LIGHT,GRAY)
line(s,10.08,4.88,11.67,4.88,BLUE,1.6,True)
label(s,9.53,5.09,3.1,"Field + memory")
group(s,start,"information-domain audit then common-scenario projection")
body(s,["Changed claim: which independent model preserves the response to the same rate change?",
        "Same public traffic, forcing and full objective; no model receives future reference answers.",
        "A sufficient compact summary is a positive outcome; hiddenness is not assumed."],y=5.62,step=.40)

s = slide("Ramp entrants create a lasting moving effect",
    "BP-R Stage 3 selected the Laval--Daganzo continuum--particle/multilane "
    "hybrid family: lane-wise Kinematic Wave (KW) conservation coupled to "
    "selected entrants' constrained speed/position trajectories. Hybrid in this "
    "source means macroscopic streams plus selected microscopic kinetics, not "
    "merely lane resolution or a discretization. Our approved adaptation "
    "keeps continuous vehicle-equivalent population and zero-population "
    "restriction metadata. Restrictions are effects of traffic already counted "
    "in the field. Upstream continuous effective-rate permission acts before "
    "finite ramp propagation; permission is not actual merge discharge or "
    "a faithfully resolved binary signal. The mandatory merge uses a fixed "
    "ramp receiving share with unused capacity available to other eligible "
    "traffic. Mainline discretionary transfers remain sampled. This physical "
    "family is literature-supported, but the whole adapted module composition "
    "and its necessity/benefit are not source theorems.",
    [F+"findings_and_reference_gate.md", F+"meter_boundary_gate.md",
     F+"merge_allocation_gate.md", F+"d2_forward/findings_and_field_gate.md",
     "https://escholarship.org/content/qt8w70q261/qt8w70q261.pdf",
     "https://bpb-us-e1.wpmucdn.com/sites.gatech.edu/dist/f/3571/files/2023/03/merges.pdf"])
lead(s,"After merging, a slow ramp entrant can continue to restrict the receiving lane.")
start=len(s.shapes)
label(s,.92,2.25,2.25,"Mainline lanes",bold=True)
rect(s,3.10,2.32,8.60,.44,ROAD)
rect(s,3.10,3.30,8.60,.56,WHITE,GRAY)
rect(s,3.13,3.33,5.13,.50,BLUE)
arrow(s,11.85,2.48,.44,.17,TEAL)
arrow(s,11.85,3.52,.44,.17,TEAL)
# Both ramp edges meet the actual lower-lane boundary; directed entry
# continues inside that lane to the downstream slow-entrant descriptor.
line(s,1.45,5.10,5.65,3.86,GRAY,2.1)
line(s,1.45,5.53,6.12,3.86,GRAY,2.1)
line(s,2.15,5.02,2.29,5.30,AMBER,3.0)
line(s,2.75,4.94,5.94,3.72,AMBER,2.5,True)
line(s,6.22,3.61,7.93,3.61,AMBER,2.5,True)
line(s,8.30,3.08,8.30,4.03,AMBER,2.7)
arrow(s,8.50,3.51,.55,.18,AMBER)
label(s,6.90,2.83,5.10,"Slow entrant after merging",color=AMBER,bold=True)
label(s,9.11,4.18,2.55,"Gap ahead")
label(s,3.20,3.37,3.45,"Slowed followers",color=WHITE)
label(s,4.58,4.75,3.85,"Actual ramp merge",color=AMBER)
line(s,5.75,4.65,5.95,3.92,AMBER,1.6,True)
label(s,1.03,5.52,3.6,"Upstream rate permission")
label(s,3.05,5.02,4.1,"Finite ramp; travel and queue")
group(s,start,"connected ramp entry becomes downstream moving restriction")
body(s,["Population stays continuous and conserved; a restriction adds no extra vehicle mass.",
        "Basis: Laval–Daganzo flow and moving effects; the joined task uses explicit adaptations."],y=6.02,step=.46)

s = slide("A mean model needs recovery memory",
    "The Stage 3 primary comparator was a mechanism-oriented present-condition "
    "mean service law; the companion was a congestion-triggered dynamic capacity "
    "drop model. Both are legitimate families, with different source domains. "
    "A's local stationary limits pass, but zero present entries is not a complete "
    "transient recovery state; B can impose a drop at sampled onset while the "
    "chosen reference has not yet created a moving restriction. These are "
    "bounded law/activation mismatches, not a proof that every aggregate model "
    "fails. T1 strengthens the primary state to distinguish startup, active "
    "effects and recovery. Transfer counter remainders supply supported event "
    "creation memory. Event times reconstruct motion only in the restricted "
    "free-ahead case; congestion changes recovery. Group counts/means lose "
    "position-speed association. M1 expands to spatial/kinetic groups; G1 "
    "selects exact distinct signatures first, without an engineering cost "
    "criterion or a concurrent reduced-model acceptance condition. Reductions "
    "are deferred until the exact reference chain is qualified.",
    [F+"dual_qualification_findings_and_scope_gate.md",
     F+"event_memory_findings_and_gate.md",F+"grouped_recovery_findings_and_gate.md",
     F+"exact_information_findings_and_gate.md"])
lead(s,"Zero new entries does not specify what old slow entrants still do.")
start=len(s.shapes)
label(s,.97,2.02,11.0,"After actual entries stop: the current entry rate is zero in both phases",bold=True)
label(s,1.03,2.74,3.50,"Soon after the last merge")
density_strip(s,4.65,2.75,5.72,1.85,"rear",BLUE)
line(s,6.56,2.55,6.56,3.48,AMBER,2.5)
line(s,6.77,2.98,7.31,2.98,AMBER,2.0,True)
label(s,4.65,3.53,7.10,"Old entrant still slows followers",color=AMBER)
line(s,1.18,3.46,1.18,4.33,GRAY,1.8,True)
label(s,1.55,3.88,2.9,"Traffic + acceleration")
label(s,1.03,4.53,3.45,"After recovery")
rect(s,4.65,4.55,5.72,.47,LIGHT,GRAY)
line(s,9.31,4.36,9.31,5.24,GRAY,1.8)
line(s,9.49,4.79,10.14,4.79,TEAL,2.0,True)
label(s,4.65,5.33,7.10,"Entrant no longer adds a service restriction",color=TEAL)
group(s,start,"two recovery phases with the same current zero entry rate")
body(s,["A mean law based on current entries can miss the continuing effect of an earlier merge.",
        "Event counters remember creation; traffic ahead determines later motion and release.",
        "We kept distinct effects first; whether a smaller summary suffices remains open."],y=5.86,step=.39)

s = slide("One traffic state must drive every component",
    "Exact-first qualification retained all restriction signatures and own "
    "history, then bound sampled insertion E1: completed actual transfers create "
    "zero-population descriptors at the regular clock, with own updated counts "
    "and donor request-phase speed, no retrospective effect. K.A resolves "
    "braking when a queue is encountered, integrates actual distance, carries "
    "the ending speed and permits positive recovery at the next regular clock. "
    "Traffic counts must adjust for accepted transfers and be rebased on actual "
    "section crossing. The alignment ladder examined unchanged interfaces, "
    "explicit restriction-aware adaptation and a count/wave representation. "
    "Uniform whole-cell fields concretely lose retained empty gaps. Step 2 "
    "supports exact local two-record evolution; neither it nor step 3 closes "
    "all endogenous transfer/motion/birth interactions. Subsequent own-field "
    "transfer work qualifies additional bounded coupled examples. This is "
    "constructive partial support, not all-mode or indefinite evolution. The "
    "diagram is schematic; no new traffic trajectory is simulated.",
    [F+"sampled_insertion_findings.md",F+"motion_contact_findings_and_gate.md",
     F+"joined_step_findings_and_gate.md",F+"alignment_ladder_findings.md",
     F+"fixed_share_completion_findings_and_gate.md"])
lead(s,"Retaining effects is useful only if their motion and traffic service share one state.")
start=len(s.shapes)
label(s,.95,2.12,3.5,"Completed spatial field",bold=True)
rect(s,4.65,2.17,7.05,.50,WHITE,GRAY)
rect(s,4.67,2.19,1.42,.46,BLUE)
rect(s,7.45,2.19,1.47,.46,BLUE)
for xx in [6.12,8.93]:line(s,xx,1.98,xx,2.91,AMBER,2.1)
label(s,6.30,2.98,2.65,"Retained empty gap",center=True)
line(s,4.05,2.95,4.05,3.63,AMBER,1.8,True)
label(s,.95,3.17,2.9,"Averaging projection",color=AMBER)
label(s,.95,3.84,3.5,"Whole-cell averaging",bold=True)
rect(s,4.65,3.89,7.05,.50,LIGHT,GRAY)
for xx in [6.12,8.93]:line(s,xx,3.70,xx,4.64,AMBER,2.1)
label(s,5.15,4.65,6.15,"Same total mass; wrong traffic between effects",color=AMBER)
group(s,start,"gap-preserving field versus cell average")
body(s,["Births and braking now have local checks; ordinary waves still govern the effects' motion.",
        "Unchanged averaging loses the gaps; restriction-aware fields give local coupled evolution.",
        "The remaining issue is a complete shared update of motion, transfers and the next field."],
     y=5.43,step=.48)

s = slide("A candidate chain; the task remains open",
    "This final primary page synthesizes the last completed D2 continuation. "
    "The whole-forward contract now consolidates upstream continuous admission, "
    "finite ramp propagation, lane-wise first-order field, own request/history "
    "sampling, T.A amount allowances, P.A held competition weights, N.A fixed "
    "ramp priority/unused receiving, D2 physical sender pooling, K.A old-record "
    "motion, accepted counts and completed E1 births. These are approved "
    "contract choices, with source/static/local support rather than whole-model "
    "adoption. The D2 static allocator has a general slack-resource invariance "
    "argument plus 5,832 exact cases, 8,292 strict donor/ramp lifts and 3,716 "
    "slack receiver lifts, all pass. This is not a traffic simulation or a "
    "whole-network existence result. A completed joint one-step field example "
    "includes quota exhaustion, wave contacts, a section crossing and two "
    "simultaneous births. The current Q.J.F.R field-carryover decision is open: "
    "A preserve completed profile (recommended), B qualify a projection. "
    "Autonomous repeated evolution, clearance, reached common-rate complete-TST "
    "leverage, independent reductions and derivative usefulness remain missing. "
    "Stage 3/67 open, Stage 4 unopened; do not infer mesoscopic necessity.",
    [F+"d2_forward/findings_and_field_gate.md",
     SR+"d2_forward/forward_contract_and_qualification.md",
     F+"fixed_share_completion_findings_and_gate.md"])
lead(s,"Local checks support the chain; useful gradients are not established.")
start=len(s.shapes)
chain(s,["Rate permission\n+ finite ramp","Lane-wise field\n↔ moving effects","Shared transfers\n+ own counts","Sampled births\n+ next state"],y=2.10)
line(s,11.24,2.96,11.24,3.20,TEAL,1.6)
line(s,11.24,3.20,5.16,3.20,TEAL,1.6)
line(s,5.16,3.20,5.16,2.96,TEAL,1.6,True)
label(s,5.37,3.26,6.5,"Completed own state → next interval",color=TEAL,h=.34)
group(s,start,"forward chain with explicit own-state update feedback")
body(s,["Established: continuous accounting, selected service rules and local coupled certificates.",
        "Static shared-service check: queue growth alone does not earn extra priority.",
        "Open now: the ordinary traffic profile preserved at the next timestep.",
        "Still needed: repeated evolution, clearance and full-system control response.",
        "Then: fair independent reductions, gradient correctness and decision usefulness."],y=3.69,step=.61)

# Four secondary pages: technical content for the small expert discussion.
s = slide("Technical: effect creation and recovery",
    "E1 attaches effects after completed actual transfers, not permission, "
    "green pulses or equal-size physical packets. An illustrative unit "
    "threshold sees 0.4 then 0.6 equivalents: no birth at the first clock, one "
    "at the second. The first amount already entered the population; its "
    "counter remainder survives pauses. Small terminal remainders still count "
    "in population/exits/TST even if no additional descriptor appears. Several "
    "simultaneous descriptors at one point get the same own count and distinct "
    "motion signatures, no artificial spacing or extra mass. The second "
    "example is the bounded moment-update gap: equal-probability speeds (0,20) "
    "and (10,10) have equal mean 10. Clipping individual speeds to an ahead "
    "limit of 10 gives means 5 and 10. Thus a mean alone cannot determine "
    "this recovery update. This is not a reached ramp-control failure or a "
    "proof against all richer compact statistics. Exact distinct records are "
    "selected before approximation; underlying moving-restriction trajectories "
    "and ordinary conserved traffic remain different parts of the state.",
    [F+"sampled_insertion_findings.md",F+"event_memory_findings_and_gate.md",
     F+"grouped_recovery_findings_and_gate.md"])
lead(s,"A transfer counter creates an effect record; it never creates extra traffic.")
start=len(s.shapes)
label(s,1.00,2.10,5.5,"Illustrative threshold: 1 equivalent",bold=True)
rect(s,1.02,2.91,8.50,.24,LIGHT)
rect(s,1.02,2.91,3.40,.24,TEAL)
rect(s,4.42,2.91,5.10,.24,BLUE)
line(s,4.42,2.76,4.42,3.47,GRAY,1.7)
line(s,9.52,2.76,9.52,3.47,AMBER,2.2)
label(s,1.02,3.57,3.6,"0.4: no new record")
label(s,5.00,3.57,6.9,"Then 0.6: one record at the next update",color=TEAL)
group(s,start,"continuous actual mass and sampled descriptor threshold")
body(s,["Counter progress persists when entries pause; all accepted mass stays accounted for.",
        "Recovery also needs the traffic ahead: identical mean speeds can evolve differently."],y=4.39,step=.55)
formula(s,"mean_braking",[
    r"(0,20)\ \mathrm{and}\ (10,10):\quad \bar v=10\ \mathrm{m/s}",
    r"v\mapsto\min(v,10):\quad \bar v_{\mathrm{after}}=5\ \mathrm{versus}\ 10\ \mathrm{m/s}"],
    1.15,5.62,11.02,1.07,fontsize=24,line_spacing=.62)

s = slide("Technical: amounts versus rate ceilings",
    "This page groups three consecutive closed methodological choices: T.A "
    "sampled requests are integrated mass allowances; P.A holds their original "
    "competition weights during the public step without a queue bonus; D2 "
    "shares the physical donor's available sending among active original "
    "requests without an additional theta-times-capacity rate ceiling. N.A "
    "still fixes the ramp's geometric share and permits unused receiving "
    "redistribution. Weights refresh on the next regular request clock; "
    "fulfilling an allowance closes that movement with no retry or enlarged "
    "amount. The late-arrival diagnostic is a 30m donor, u=30m/s, 1s step, "
    "n=.45eq, sampled lateral fraction .25 and Q=9/14eq/s. Traffic in the "
    "first 3m reaches the exit at .9s; through receiving is blocked and lateral "
    "receiving free. D1's extra .25Q rate ceiling would serve 9/560=.016071eq; "
    "approved D2 serves 9/140=.064286eq, both below d=.1125eq. These are "
    "distinct timing laws for the same request, not a hidden-state witness. "
    "D1 remains conditionally sound if an independently justified persistent "
    "direction-specific instantaneous cap exists. Current evidence does not "
    "require it. D2's fixed-weight static queue-invariance result is positive "
    "only for the finite allocator, not the whole endogenous evolution.",
    [F+"transfer_feedback_findings_and_gate.md",F+"quota_allocation_findings_and_gate.md",
     F+"fixed_share_completion_findings_and_gate.md",F+"d2_forward/findings_and_field_gate.md"])
lead(s,"A 25% lane-change request limits mass; it need not limit instantaneous sending.")
formula(s,"amount_and_service",[
    r"d_{\mathrm{lat}}=0.25\times0.45=0.1125\ \mathrm{eq},\qquad Q=0.6429\ \mathrm{eq/s}",
    r"\Delta t_{\mathrm{available}}=0.1\ \mathrm{s}:\quad 0.25Q\Delta t=0.0161\ \mathrm{eq},\quad Q\Delta t=0.0643\ \mathrm{eq}"],
    .95,2.00,11.48,1.22,fontsize=24,line_spacing=.64)
start=len(s.shapes)
label(s,1.03,3.85,4.6,"Extra direction-specific rate cap")
label(s,1.03,4.75,4.98,"Accepted: shared sending",color=TEAL,bold=True)
for yy,width,color in [(3.94,1.18,GRAY),(4.84,4.72,TEAL)]:
    rect(s,6.10,yy,4.72,.28,LIGHT)
    rect(s,6.10,yy,width,.28,color)
label(s,6.10,5.40,5.7,"Accepted mass in the final 0.1 s")
group(s,start,"same late arrival under two explicit sender rules")
body(s,["Through is blocked; the lateral target is free; the amount allowance is not exhausted.",
        "Queue growth changes availability, not the fixed competition weight or allowed mass."],y=6.06,step=.51)

s = slide("Technical: records do not preserve the field",
    "The latest completed D2 pass supplies this ordinary-field diagnostic. "
    "Both 30m donors have mass .45eq, the same fixed-end counts and no "
    "restriction records. Their common sampled lateral allowance is .1125eq, "
    "through receiving is blocked, lateral receiving is free, and Q=9/14eq/s. "
    "The early field has jam density .15eq/m on the last 3m [27,30]. Sending "
    "starts immediately and the allowance is exhausted after .175s, giving "
    ".1125eq. The late field has the same jam block on [0,3]; its front "
    "reaches the exit at .9s and serves Q*.1=9/140=.064286eq. These are "
    "admissible first-order field/request diagnostics, not two certified "
    "reachable equal-history scenarios. A mass-plus-anchor reconstruction "
    "maps them to the same field, so conservation and exact retained records "
    "do not imply ordinary profile/timing preservation. A sufficiently "
    "resolved macroscopic field can distinguish them. No external exit before "
    "the end means one-step population-time is the same; no full-TST benefit "
    "or metadata necessity follows. Current Q.J.F.R remains open: A retain "
    "completed full profile is recommended; B qualify a finite numerical "
    "mass-and-anchor reconstruction is legitimate if its error/reference "
    "contract is justified. Neither selected by this presentation.",
    [F+"d2_forward/findings_and_field_gate.md",SR+"d2_forward/forward_math.tex"])
lead(s,"Where the same mass sits determines when it can reach the transfer port.")
start=len(s.shapes)
label(s,1.04,2.12,10.9,"Same 30 m cell; 0.45 equivalents; same request and boundary counts",bold=True)
label(s,1.04,3.02,3.15,"Traffic at the exit")
density_strip(s,4.34,3.03,6.55,.655,"front",TEAL)
label(s,4.34,3.75,7.36,"Immediate service; 0.1125 equivalents transferred",color=TEAL)
label(s,1.04,4.58,3.15,"Traffic at the entrance")
density_strip(s,4.34,4.59,6.55,.655,"rear",BLUE)
line(s,5.18,5.38,10.88,5.38,BLUE,1.8,True)
label(s,4.34,5.66,7.70,"Arrival at 0.9 s; 0.0643 equivalents transferred",color=BLUE)
group(s,start,"same cell totals different physical arrival times")
body(s,["A spatial macro model can retain this difference; it is not an identity-necessity result.",
        "Open choice: carry the completed profile, or qualify an explicit numerical reconstruction."],y=6.10,step=.44)

s = slide("Technical: the useful-gradient evidence chain",
    "Complete all-stream total system time J is the integral until clearance "
    "of the whole not-yet-exited population: outside waiting, the finite ramp "
    "and every modeled mainline lane. Vehicle-equivalent seconds are the "
    "units. Mean time is J divided by the fixed total cohort mass if needed; "
    "this slide uses total time. Relocating waiting from mainline to ramp or "
    "outside does not improve this objective unless total exit timing changes. "
    "The population curves drawn here are conceptual layers, not calculated "
    "trajectories. The general sensitivity path is rate permission to actual "
    "admission/travel/queue, accepted transfers and moving effects, exits, "
    "then complete J. Structural topology, record membership and creation "
    "events require a declared derivative/event contract, not an unsupported "
    "global smoothness claim. Exact forward qualification, clearance and a "
    "reached intervention-sensitive response come before specific gradients. "
    "Strong independent spatial/memory/moment rivals and declared initial "
    "projections make the representation comparison fair. AD/finite-difference "
    "correctness and independent-reference directional usefulness are later "
    "distinct evidence products. No optimizer, SUMO, traffic simulation or "
    "numerical gradient was performed for this review or presentation. This "
    "is the discussion endpoint, not an automatic Stage 4 opening.",
    [SR+"d2_forward/forward_contract_and_qualification.md",
     F+"baseline_review_scope_check.md",R+"/rm12_30_provisional_four_stage_roadmap.md"])
lead(s,"A useful direction must improve complete exit timing, including traffic held back.")
start=len(s.shapes)
for y,w,col,caption in [(2.19,4.35,GRAY,"Waiting outside"),
                         (2.94,3.35,AMBER,"On the ramp"),
                         (3.69,6.25,BLUE,"On the mainline")]:
    label(s,1.03,y,3.1,caption)
    rect(s,4.45,y+.02,w,.31,col)
label(s,4.45,4.26,7.3,"Count every not-yet-exited stream (schematic)")
group(s,start,"all-stream population included in system time")
formula(s,"complete_tst",[
    r"J=\int_0^{t_{\mathrm{clear}}}\!\left[N_{\mathrm{outside}}(t)+N_{\mathrm{ramp}}(t)+N_{\mathrm{main}}(t)\right]\,\mathrm{d}t"],
    .98,4.92,11.44,.72,fontsize=25)
body(s,["N: remaining vehicle-equivalent mass; J: total system time until every stream clears.",
        "Finish the forward chain and fair reductions, then verify derivatives and decision gains."],y=6.02,step=.46)


ARGUMENTS = {
    3: {"purpose": "Set the common research chain against which every later result is judged.",
        "incoming": "Last week established that equal mass alone does not make lane information unavailable.",
        "finding": "Physical support, unresolved information, useful common actions and useful gradients are distinct obligations.",
        "outgoing": "The original history claim motivates a search for genuine unresolved traffic detail."},
    4: {"purpose": "Use last week's obstruction to explain why this week's mechanism reviews were needed.",
        "incoming": "The original question asks for information unavailable after all declared observations and history.",
        "finding": "Flow can reveal lane allocation, and known complete initialization plus exact public dynamics can determine later states; these are conditional results.",
        "outgoing": "Therefore we tested recovery and finite merge space before changing the claim."},
    5: {"purpose": "Name the proposed missing information and show why physical mechanism support did not certify hiddenness.",
        "incoming": "The elementary lane-split construction did not establish a complete-history residual.",
        "finding": "Entrant location/speed/recovery and lane-end distance/gap timing are supported physical leads, but the reviewed sources did not join full-history ambiguity to useful common-rate actions.",
        "outgoing": "Because the full-history link remained missing, we reconsidered that claim."},
    6: {"purpose": "Explain the actual backtracking point: from unresolved knowledge to loss caused by a declared model reduction.",
        "incoming": "Two physical extensions failed the same history/action admission link; operational detector inputs did not bind all initial/history knowledge either.",
        "finding": "BP-R permits common known physical initialization with declared model projections and independent evolution; it does not establish the original hidden-history claim.",
        "outgoing": "This change lets us audit a shared physical model before asking what can be removed."},
    7: {"purpose": "Bind the supported physical mechanism to the unchanged upstream metering task.",
        "incoming": "A model-reduction comparison needs one specified physical reference and common control authority.",
        "finding": "The connected ramp arrow explicitly traces a transferred entrant into a downstream moving restriction; lane-wise conserved flow and selected entrant kinetics are coupled.",
        "outgoing": "The next question is what state remembers the entrant's continuing traffic effect."},
    8: {"purpose": "Explain why current-entry mean service was not a complete transient recovery model, and why exact-state qualification was selected first.",
        "incoming": "Slow entrants can continue affecting service after the upstream action and actual merge entries have changed.",
        "finding": "Zero current entries is compatible with an active old restriction or recovered motion. By itself that signal cannot select the correct extra service effect; event counters certify creation, not all recovery.",
        "outgoing": "We retained exact effects to qualify the chain before testing smaller summaries."},
    9: {"purpose": "Show why retaining exact disturbance descriptors did not automatically complete the forward model.",
        "incoming": "Exact-first retention removes the need to invent the discarded recovery arrangement.",
        "finding": "Motion needs the same traffic ahead that service updates; averaging a field can erase retained empty gaps even while mass and records survive. Local richer-field evolution is supported, whole composition remains partial.",
        "outgoing": "Information value can only be tested after the shared update predicts its own next state."},
    10: {"purpose": "Separate the established model contract and local evidence from the unestablished whole-task research claims.",
        "incoming": "Birth, motion, transfer and field audits provide a recognizable common model chain.",
        "finding": "Selected rules and bounded local/static checks are consolidated; autonomous repetition, clearance, control response and fair reductions are not completed.",
        "outgoing": "Local construction is supported; full control-response and gradient claims remain open."},
    11: {"purpose": "Explain the two retained-state duties: actual-transfer creation bookkeeping and traffic-dependent recovery.",
        "incoming": "The main presentation selected exact-state qualification before reduction.",
        "finding": "A threshold counter creates no population, while equal means can evolve differently under individual braking; the latter is a local update diagnostic.",
        "outgoing": "Creation memory is qualified locally; recovery still needs the surrounding traffic field."},
    12: {"purpose": "Explain how the accepted amount/weight/sending choices help complete the shared conservative transfer interface.",
        "incoming": "A moving restriction model needs a port rule consistent with actual sending, receiving and remaining movement requests.",
        "finding": "The accepted D2 completion imposes no extra oriented instantaneous rate cap; original amount ceilings and held competition weights remain.",
        "outgoing": "The accepted service rule helps join disturbance state to conservative traffic flow."},
    13: {"purpose": "Identify the last completed checkpoint's state contract and why it must close before a fair reduction comparison.",
        "incoming": "D2's static map is qualified, but each next interval also needs its ordinary spatial traffic field.",
        "finding": "Equal cell mass and fixed-end count anchors can yield different arrival timing. This is an ordinary spatial-profile diagnostic, not a reached-history or mesoscopic-necessity certificate.",
        "outgoing": "A faithful reference is needed before errors can be attributed to omitted effect records."},
    14: {"purpose": "Reconnect the unfinished forward-model work to the ultimate control and gradient research goal.",
        "incoming": "A local flow prediction difference is weaker than complete intervention-sensitive exits and total system time.",
        "finding": "All waiting populations must count; correct derivatives and independently useful directions require different evidence.",
        "outgoing": "A correct forward model and a useful control direction require separate evidence."},
}

for number, argument in ARGUMENTS.items():
    current = prs.slides[number-1]
    text(current, .80, 7.03, 11.88, .34, argument["outgoing"],
         role="argument", color=TEAL)
    record = records[number-1]
    record["argument"] = argument
    record["notes"] = ("Presentation purpose: "+argument["purpose"]+"\n\n"
        "Connection from the preceding question: "+argument["incoming"]+"\n\n"
        "Finding and scope: "+argument["finding"]+"\n\n"
        "Why this leads onward: "+argument["outgoing"]+"\n\n"+record["notes"])
    if number == 4:
        record["notes"] += ("\n\nAmendment: the earlier normalized mass/flow "
            "calculation is no longer displayed; this page now explains only "
            "the consequence that opened this week's review.")
    if number == 8:
        record["notes"] += ("\n\nThe two phases are schematic recovery "
            "conditions, not a new reached trajectory or equal-aggregate-history "
            "pair. Their only deliberately shared signal is zero current actual "
            "entry rate. The point concerns the present-entry penalty component, "
            "not a refutation of all spatial, memory or moment models. A record "
            "can remain retained when no longer binding; no automatic deletion "
            "or new recovery law is introduced by this illustration.")
    current.notes_slide.notes_text_frame.text = (record["notes"]+
        "\n\nSources:\n"+"\n".join(record["sources"]))



def leaves(shapes, ox=0, oy=0, sx=1, sy=1):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xf = shape._element.grpSpPr.xfrm
            co, ce = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
            gx, gy = shape.width/int(ce.get("cx")), shape.height/int(ce.get("cy"))
            yield from leaves(shape.shapes,
                              ox+sx*(shape.left-int(co.get("x"))*gx),
                              oy+sy*(shape.top-int(co.get("y"))*gy),sx*gx,sy*gy)
        else:
            yield shape, ((ox+sx*shape.left)/EMU,(oy+sy*shape.top)/EMU,
                          sx*shape.width/EMU,sy*shape.height/EMU)


def font(points,bold=False,scale=144):
    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans"
                              +("-Bold" if bold else "")+".ttf",round(points*scale/72))


def validate(saved):
    assert len(saved.slides)==14 and sum(r["boilerplate"] for r in records)==2
    counts=[]; geometric=[]; visible=[]
    for i,s in enumerate(saved.slides):
        lines=0; equations=0; titles=0
        assert s.background.fill.fore_color.rgb==rgb(WHITE)
        geometric.append(sum(sh.shape_type==MSO_SHAPE_TYPE.GROUP for sh in s.shapes))
        for sh,(x,y,w,h) in leaves(s.shapes):
            assert min(x,y)>=-.01 and x+w<=13.344 and y+h<=7.51,(i+1,sh.name,"off slide")
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE and sh.name.startswith("math:"):
                equations+=math_lines[sh.name]
            if not sh.has_text_frame or not sh.text:
                continue
            visible.append(sh.text)
            istitle=sh.name.startswith("title:")
            titles+=int(istitle)
            if not istitle and not sh.name.startswith("plot:"):
                lines+=len(sh.text.splitlines())
            tf=sh.text_frame
            for p in tf.paragraphs:
                pts=32 if istitle else 18
                assert p.font.name==FONT and p.font.size.pt==pts,(i+1,sh.name,"font")
                avail=(w-(tf.margin_left+tf.margin_right)/EMU)*144
                assert font(pts,bool(p.font.bold)).getlength(p.text)<=avail,(i+1,sh.name,"text width")
            required=len(sh.text.splitlines())*(32 if istitle else 18)*1.15/72
            availh=h-(tf.margin_top+tf.margin_bottom)/EMU
            assert required<=availh+.04,(i+1,sh.name,"text height")
        assert titles==1 and lines<=8 and lines+equations<=8,(i+1,lines,equations)
        assert len(s.notes_slide.notes_text_frame.text)>250
        counts.append({"slide":i+1,"body_lines":lines,"formula_lines":equations})
    assert not re.search(r"\b(S0[1-9]|S1[0-3]|RM[0-9]|UG[1-4]|Candidate [ABC]|M2C|Phase [A-F])\b","\n".join(visible))
    for record in records:
        for source in record["sources"]:
            if "://" not in source:
                assert (ROOT/source).exists(),source
    with zipfile.ZipFile(OUT) as archive:
        assert archive.testzip() is None
        names=set(archive.namelist())
        for member in names:
            if member.endswith(".xml") or member.endswith(".rels") or member.endswith(".svg"):
                etree.fromstring(archive.read(member))
        vector_count=sum(name.endswith(".svg") for name in names)
        assert vector_count==3
        svg_refs=0
        for i in range(1,15):
            xml=etree.fromstring(archive.read(f"ppt/slides/slide{i}.xml"))
            rels=etree.fromstring(archive.read(f"ppt/slides/_rels/slide{i}.xml.rels"))
            relmap={node.get("Id"):node.get("Target") for node in rels}
            for node in xml.findall(".//{"+SVG_NS+"}svgBlip"):
                rid=node.get(qn("r:embed"))
                assert rid in relmap
                target=posixpath.normpath(posixpath.join("ppt/slides",relmap[rid]))
                assert target in names and target.endswith(".svg")
                svg_refs+=1
        assert svg_refs==3
    return {"slides":14,"boilerplate":2,"substantive":12,"body_line_counts":counts,
            "grouped_diagrams_by_slide":geometric,"embedded_vector_assets":vector_count,
            "font":"Calibri","title_points":32,"body_points":18,
            "white_background":True,"xml_zip_and_relationship_checks":"PASS",
            "source_paths":"PASS","layout_checks":"PASS",
            "argument_links": {str(n): a for n,a in ARGUMENTS.items()},
            "preview":"Pillow layout proof using DejaVu Sans; not Office rendering"}


def color_of(value,default):
    try:return "#"+str(value.rgb)
    except (AttributeError,ValueError,TypeError):return default


def preview(saved):
    scale=120; pages=[]
    for i,s in enumerate(saved.slides,1):
        im=Image.new("RGB",(1600,900),"#"+WHITE)
        draw=ImageDraw.Draw(im)
        for sh,(x,y,w,h) in leaves(s.shapes):
            x,y,w,h=(v*scale for v in (x,y,w,h))
            bounds=(x,y,x+w,y+h)
            if sh.shape_type==MSO_SHAPE_TYPE.LINE:
                xf=sh._element.spPr.xfrm
                xx=(x+w,x) if xf.get("flipH")=="1" else (x,x+w)
                yy=(y+h,y) if xf.get("flipV")=="1" else (y,y+h)
                col=color_of(sh.line.color,"#"+GRAY)
                draw.line((xx[0],yy[0],xx[1],yy[1]),fill=col,width=3)
                if sh._element.spPr.ln.find(qn("a:tailEnd")) is not None:
                    angle=math.atan2(yy[1]-yy[0],xx[1]-xx[0]);tip=(xx[1],yy[1])
                    draw.polygon([tip,(tip[0]-13*math.cos(angle-.35),tip[1]-13*math.sin(angle-.35)),
                                  (tip[0]-13*math.cos(angle+.35),tip[1]-13*math.sin(angle+.35))],fill=col)
                continue
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                picture=Image.open(BytesIO(sh.image.blob)).convert("RGBA")
                picture=picture.resize((max(1,round(w)),max(1,round(h))),Image.Resampling.LANCZOS)
                im.paste(picture,(round(x),round(y)),picture)
                continue
            if sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill=color_of(sh.fill.fore_color,None) if sh.fill.type==1 else None
                col=color_of(sh.line.color,None)
                kind=sh.auto_shape_type
                if kind==MSO_SHAPE.OVAL:
                    draw.ellipse(bounds,fill=fill,outline=col,width=2)
                elif kind in [MSO_SHAPE.RIGHT_ARROW,MSO_SHAPE.UP_ARROW]:
                    if kind==MSO_SHAPE.RIGHT_ARROW:
                        points=[(x,y+h*.25),(x+w*.62,y+h*.25),(x+w*.62,y),
                                (x+w,y+h*.5),(x+w*.62,y+h),(x+w*.62,y+h*.75),(x,y+h*.75)]
                    else:
                        points=[(x+w*.25,y+h),(x+w*.25,y+h*.38),(x,y+h*.38),
                                (x+w*.5,y),(x+w,y+h*.38),(x+w*.75,y+h*.38),(x+w*.75,y+h)]
                    draw.polygon(points,fill=fill,outline=col)
                else:
                    draw.rectangle(bounds,fill=fill,outline=col,width=2)
            if not sh.has_text_frame or not sh.text:continue
            tf=sh.text_frame
            ml,mt=tf.margin_left/EMU*scale,tf.margin_top/EMU*scale
            mr,mb=tf.margin_right/EMU*scale,tf.margin_bottom/EMU*scale
            paragraphs=list(tf.paragraphs);pts=paragraphs[0].font.size.pt
            step=pts*scale/72*1.15;ty=y+mt
            if tf.vertical_anchor==MSO_ANCHOR.MIDDLE:
                ty=y+mt+(h-mt-mb-len(paragraphs)*step)/2
            for p in paragraphs:
                f=font(p.font.size.pt,bool(p.font.bold),scale)
                tx=x+ml
                if p.alignment==PP_ALIGN.CENTER:
                    tx+=(w-ml-mr-f.getlength(p.text))/2
                draw.text((tx,ty),p.text,font=f,fill=color_of(p.font.color,"#"+INK),anchor="lt")
                ty+=step
        im.save(ASSETS/f"slide_{i:02d}_layout.png")
        pages.append(im)
    contact=Image.new("RGB",(1366,7*401+8),"#DCE1E5")
    for i,im in enumerate(pages):
        thumb=im.resize((667,375),Image.Resampling.LANCZOS)
        contact.paste(thumb,(8+(i%2)*683,8+(i//2)*401))
    contact.save(PREVIEW)
    for start in [0,8]:
        subset=pages[start:start+8]; rows=math.ceil(len(subset)/2)
        sheet=Image.new("RGB",(1366,rows*401+8),"#DCE1E5")
        for j,im in enumerate(subset):
            sheet.paste(im.resize((667,375),Image.Resampling.LANCZOS),
                        (8+(j%2)*683,8+(j//2)*401))
        sheet.save(ASSETS/f"contact_{start+1:02d}_{start+len(subset):02d}.png")


def write_notes(result):
    content = [
        "# Weekly ramp-metering task review — request 7", "",
        "Date: 2026-10-09. Continuation of last week's mass/flow information-hiding "
        "obstruction through the approved D2 choice and its completed dependent "
        "qualification. Two boilerplate slides and twelve substantive slides. "
        "The review is temporarily parked at the still-open ordinary-field carryover gate. "
        "Revised for the request-7 narrative/diagram amendment.", "",
        "## The argument connecting the slides", "",
        "Last week's flow/mass result meant that a density-equal snapshot "
        "did not establish useful hidden information. Recovery and finite "
        "merge geometry supplied genuine physical leads, but neither joined "
        "that detail to full-history ambiguity and a useful common rate action. "
        "An operational sensing audit also left total initial/history access "
        "unbound. The explicit backtracking therefore changed the claim to "
        "what a declared model loses when it reduces shared known traffic.", "",
        "Under that revised claim, the supported moving-effect family needs "
        "a model that remembers recovery after actual entries stop. Exact "
        "records avoid inventing an unqualified mean arrangement, but they "
        "must still evolve with one compatible ordinary traffic field. "
        "Local motion/birth/transfer and static service results support that "
        "construction; repeated evolution, clearance and full-system rate "
        "response remain missing. Only then can fair reductions and separate "
        "gradient evidence answer the broader research goal.", "",
        "Every substantive slide now states a finding and a causal takeaway "
        "connecting it to that argument. Its PowerPoint notes explicitly "
        "record purpose, preceding question, finding and next implication. "
        "The original deck/generator/notes are preserved under "
        "weekly_freeway_control_review_20261009_assets/before_amendment/.", "",
        "## How to present the deck", "",
        "Use slides **3–10** as the group-meeting backbone. They tell the whole "
        "review story without requiring internal checkpoint vocabulary. Use "
        "slides **11–14** for the smaller technical conversation afterward. "
        "Every slide has detailed PowerPoint speaker notes and traceable sources; "
        "those notes are also reproduced below.", "",
        "| Slides | Purpose |",
        "|---|---|",
        "| 3–4 | State the research chain; use last week's obstruction only to motivate this week's search |",
        "| 5 | Recovery and finite geometry: physical support, scoped negative admission |",
        "| 6 | Operational-information audit, then explicit shift to model-reduction adequacy |",
        "| 7–9 | Supported slow-entrant family, memory limits, exact records and common-field coupling |",
        "| 10 | Consolidated stack and unfinished whole-task claims |",
        "| 11–13 | Creation/recovery, amount-versus-rate and spatial-arrival examples |",
        "| 14 | Complete objective and the separate future gradient-evidence obligations |", "",
        "## Evidence and interpretation", "",
        "The original target concerned traffic unresolved after all admitted "
        "aggregate history. The approved backtracking asks what an independently "
        "evolving model loses by compressing a shared known physical scenario. "
        "These are distinct claims. A sufficiently informative compact model is "
        "a favorable outcome, and no macroscopic family is assumed lane blind.", "",
        "Recovery means a slow entrant's evolving acceleration, motion and "
        "restriction of following traffic. The selected source's multilane "
        "hybrid combines ordinary lane-wise KW conservation with selected "
        "constrained entrant trajectories. Continuous traffic population and "
        "zero-population restriction metadata are separate in the current "
        "adaptation. A cell grid is a numerical representation, not the "
        "physical reason this source is called hybrid.", "",
        "The slide examples reproduce existing normalized algebra, local "
        "moment/counter diagnostics and the reviewed one-step flow calculations. "
        "They are not new scientific scenarios or traffic experiments. Equal "
        "cell totals or different transfer amounts do not prove equal reachable "
        "histories, necessity of metadata, a full-system control benefit or "
        "a useful gradient. No simulator, SUMO, optimizer or numerical traffic "
        "gradient was executed.", "",
        "## Checkpoint coverage for subsequent discussion", "",
        "| Review segment | Outcome and location in deck |",
        "|---|---|",
        "| RM9 observation/history continuation | TST-related speed support; known complete-state deterministic sufficiency remains conditional (3–4). |",
        "| RM10 recovery/acceleration | Mechanism supported, full-history/common-action admission not established; scoped negative closed (5). |",
        "| RM11 finite-lane opportunity | Genuine geometry and strong spatial rivals; residual/history/action chain missing; scoped negative closed (5). |",
        "| RM12 BP-I | Used detector inputs differ from an ideal observer, but total initial/history access unbound (6). |",
        "| RM12 BP-R | Explicitly revised representation/comparison claim; source-family support motivates bounded formulation (6–7). |",
        "| Reference, upstream meter and mandatory merge choices | Selected continuum–particle family; effective-rate admission to finite ramp; fixed receiving priority and unused-share redistribution (7). |",
        "| A-primary/B-companion and transient memory | Mean-service and capacity-drop families; startup/recovery mismatches bounded; creation counter supported (8, 11). |",
        "| Grouped spatial/kinetic state and exact-first G1 | Identical-signature grouping conditionally safe; arbitrary averaging unqualified; preserve distinct records before reduction (8–9, 11). |",
        "| E1 insertion and K.A contact timing | Sampled births from actual completed transfers; no extra mass; true path, braking and ending-speed distinction (9, 11). |",
        "| Field alignment ladder | Unchanged uniform field negative; restriction-aware adaptation strongest partial construction; count/wave replacement no full success (9). |",
        "| T.A, P.A, D2 and joined transfers | Amount ceilings, held weights and shared sender; local waves/crossing/birth certificate, static queue-invariance qualification (10, 12). |",
        "| D2 dependent continuation | Whole-forward contract consolidated; exact records/count anchors do not preserve ordinary arrival profile; Q.J.F.R remains open (10, 13). |",
        "| Remaining Stage 3 / later Stage 4 | Autonomous repetition, clearance, reached complete-TST response and fair reductions before derivative correctness/usefulness (10, 14). |", "",
        "## Artifact work report", "",
        "Files created: the PPTX, this notes/work report, the self-contained "
        "Python generator, validation JSON, contact preview and asset folder "
        "(editable native diagram groups; three scalable equation assets with "
        "PNG fallbacks; per-slide layout previews). Review-only updates: README "
        "park banner, dialogue detour record and parked handoff. The feedback "
        "inbox request 7 receives only its allowed '(answered)' marker. Last "
        "week's deck/generator/assets and all scientific evidence are preserved. "
        "This amendment changes only the presentation, its generation records, "
        "the allowed feedback marker and a concise dialogue entry; it does not "
        "advance the parked review.", "",
        "Amendment changes: shorten slide 4 to the causal starting obstacle; "
        "name recovery and merge-opportunity detail on slide 5; show the "
        "actual ramp-to-lane path and downstream effect on slide 7; replace "
        "slide 8's abstract duration bars with two explicitly contrasted "
        "recovery phases sharing zero current actual entries; add directed "
        "dependency/projection arrows and argument takeaways throughout. "
        "The previous mass/flow formula is not repeated in the visible deck. "
        "All source and decision boundaries are unchanged.", "",
        "Assumptions: retain last week's white/Calibri 32/18 styling and 16:9 "
        "geometry; use twelve substantive pages and separate audiences by "
        "ordering rather than extra boilerplate or an extra appendix. Presentation "
        "synthesis neither opens review work nor selects a pending candidate.", "",
        "Build command:", "",
        "    /home/zpz/miniconda3/envs/differential_sim/bin/python misc/slides/generate_weekly_freeway_control_review_20261009.py", "",
        "Existing packages: python-pptx 1.0.2, matplotlib 3.10.9, Pillow 12.2.0, "
        "lxml and NumPy. Read the previous generator and reviewed reports; reuse "
        "its inspected native-shape/SVG relationship interfaces. No installation "
        "or production project changes.", "",
        "Validation: saved-PPTX reload; exact slide and boilerplate counts; "
        "white background; Calibri title/body sizes; body/formula line limits; "
        "text width/height and slide boundaries; one group per self-created "
        "diagram; three embedded vectors and PNG fallbacks; ZIP/XML and "
        "relationship integrity; source-file existence. Detailed result:", "",
        "[Validation JSON](weekly_freeway_control_review_20261009_validation.json)", "",
        "Rendering limitation: Calibri and Office/LibreOffice rendering are "
        "unavailable locally. The previews inspect saved-PPTX geometry and "
        "fallback pictures with DejaVu Sans, and are explicitly a layout proof, "
        "not a native Office rendering. Visual inspection covers all slides "
        "and the technical calculations; the deck remains editable in PowerPoint.", "",
        "Scientific status: review parked, Stage 3/67 open, D2 approved, "
        "Q.J.F.R undecided, Stage 4 and experiments unopened. The next allowed "
        "review action on return is discussion/selection of the field "
        "carryover contract, followed by its authorized qualification.", "",
        "## Slide-by-slide speaker notes", ""
    ]
    for record in records:
        audience = "Boilerplate" if record["boilerplate"] else (
            "Group meeting" if record["slide"] <= 10 else "Technical discussion")
        content += [f"### {record['slide']}. {record['title']}", "", audience,
                    "", record["notes"], "", "Sources:", ""]
        for source in record["sources"]:
            if "://" in source:
                content.append("- [Reviewed primary source]("+source+")")
            else:
                content.append("- ["+source+"]("+os.path.relpath(ROOT/source,HERE)+")")
        content.append("")
    NOTES.write_text("\n".join(content)+"\n")


if __name__ == "__main__":
    prs.save(OUT)
    saved = Presentation(OUT)
    result = validate(saved)
    result["presentation_route"] = {"group_meeting": list(range(3,11)),
                                    "technical_discussion": list(range(11,15))}
    result["scientific_execution"] = "None; existing source/static results synthesized only"
    result["pending_review_gate"] = "Q.J.F.R; neither candidate selected"
    preview(saved)
    write_notes(result)
    REPORT.write_text(json.dumps(result, indent=2)+"\n")
    print(json.dumps({key:result[key] for key in ["slides","boilerplate",
        "substantive","embedded_vector_assets","xml_zip_and_relationship_checks",
        "source_paths","layout_checks","pending_review_gate"]}, indent=2))
    print(OUT)
    print(PREVIEW)


