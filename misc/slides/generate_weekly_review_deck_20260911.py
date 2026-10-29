"""Generate the colleague-facing weekly review deck for 11 September 2026."""

from pathlib import Path
import math

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUT = HERE / "weekly_reviews_and_downstream_task_20260911.pptx"
FORMULA_DIR = HERE / "weekly_reviews_20260911_formulae"
FORMULA_DIR.mkdir(exist_ok=True)

FONT = "Calibri"
NAVY = RGBColor(25, 43, 64)
TEAL = RGBColor(29, 132, 139)
BLUE = RGBColor(64, 105, 156)
GREEN = RGBColor(56, 139, 92)
AMBER = RGBColor(204, 139, 40)
RED = RGBColor(174, 76, 69)
INK = RGBColor(35, 43, 50)
MID = RGBColor(93, 108, 121)
PALE = RGBColor(233, 243, 245)
LIGHT = RGBColor(247, 249, 250)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(.06)
    frame.margin_top = frame.margin_bottom = Inches(.03)
    p = frame.paragraphs[0]; p.text = value; p.alignment = align
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, value, fill=LIGHT, line=TEAL, size=18,
        bold=False, align=PP_ALIGN.CENTER):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(1.2)
    f = s.text_frame; f.clear(); f.word_wrap = True; f.vertical_anchor = MSO_ANCHOR.MIDDLE
    f.margin_left = f.margin_right = Inches(.10)
    p = f.paragraphs[0]; p.text = value; p.alignment = align
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = INK
    return s


def title(slide, value, kicker=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    text(slide, .55, .29, 12.2, .55, value, 32, NAVY, True)
    if kicker:
        text(slide, .58, .88, 12.0, .3, kicker, 14, MID)


def bullets(slide, x, y, w, h, items, size=18, color=INK):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = s.text_frame; f.clear(); f.word_wrap = True
    f.margin_left = Inches(.08); f.margin_right = Inches(.03)
    for i, item in enumerate(items):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.text = item; p.level = 0; p.font.name = FONT; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(7); p.bullet = True
        for run in p.runs:
            run.font.name = FONT; run.font.size = Pt(size); run.font.color.rgb = color
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=1.7):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width); c.line.end_arrowhead = True
    return c


def formula(name, lines, width=11.8, fontsize=22):
    path = FORMULA_DIR / f"{name}.png"
    fig = plt.figure(figsize=(width, .72 * len(lines) + .28), dpi=180)
    fig.patch.set_alpha(0)
    for i, line in enumerate(lines):
        fig.text(.5, 1 - (i + .72) / len(lines), f"${line}$", ha="center", va="center",
                 fontsize=fontsize, color="#232b32")
    plt.axis("off"); plt.savefig(path, transparent=True, bbox_inches="tight", pad_inches=.05); plt.close(fig)
    return path


def picture(slide, path, x, y, w, h):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w-pw)/2), Inches(y + (h-ph)/2), Inches(pw), Inches(ph))


def pill(slide, x, y, w, value, fill=PALE, line=TEAL, size=16):
    return box(slide, x, y, w, .48, value, fill, line, size, True)


# 1 — title boilerplate
s = prs.slides.add_slide(blank)
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.72), Inches(1.15), Inches(.15), Inches(4.7))
accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
text(s, 1.2, 1.35, 10.9, 1.35, "From Selective Discreteness\nto a Testable Control Task", 32, WHITE, True)
text(s, 1.22, 3.05, 10.5, .7, "Literature positioning · baseline design · green-allocation derivative contract", 21, RGBColor(205, 221, 228))
text(s, 1.22, 5.65, 10.6, .4, "Weekly research review | 11 September 2026", 18, WHITE)

# 2 — contents boilerplate
s = prs.slides.add_slide(blank); title(s, "Contents")
for i, (n, label) in enumerate([
    ("01", "Why task-first changes the research workflow"),
    ("02", "Review 1 — novelty, gap, and baseline ladder"),
    ("03", "Review 2 — downstream task and controlled variants"),
    ("04", "Derivative paths and current decision boundary"),
]):
    y = 1.35 + i * 1.05
    box(s, .85, y, 1.05, .68, n, TEAL, TEAL, 18, True)
    box(s, 2.12, y, 10.1, .68, label, LIGHT, PALE, 18, False, PP_ALIGN.LEFT)

# 3
s = prs.slides.add_slide(blank); title(s, "Why anchor the model to a downstream task?")
text(s, .75, 1.14, 5.25, .35, "MODEL-FIRST", 18, RED, True, PP_ALIGN.CENTER)
box(s, .65, 1.68, 1.55, .9, "Method", PALE, RED, 18, True)
box(s, 2.55, 1.68, 1.55, .9, "Build", PALE, RED, 18, True)
box(s, 4.45, 1.68, 1.55, .9, "Find task", PALE, RED, 18, True)
arrow(s, 2.2, 2.13, 2.53, 2.13, RED); arrow(s, 4.1, 2.13, 4.43, 2.13, RED)
bullets(s, .8, 3.02, 5.15, 2.55, [
    "Risk: retain detail without knowing whether it changes a decision.",
    "Risk: a convenient task may not isolate the proposed mechanism.",
], 18)
text(s, 7.25, 1.14, 5.25, .35, "TASK-ANCHORED", 18, GREEN, True, PP_ALIGN.CENTER)
box(s, 8.7, 1.55, 2.2, .82, "Methodological\nmodel", PALE, GREEN, 18, True)
box(s, 7.0, 3.02, 2.2, .82, "Implementation", LIGHT, GREEN, 18, True)
box(s, 10.4, 3.02, 2.2, .82, "Task + evidence", LIGHT, GREEN, 18, True)
arrow(s, 9.2, 2.38, 8.25, 3.0, GREEN); arrow(s, 10.4, 3.43, 9.22, 3.43, GREEN); arrow(s, 10.37, 3.43, 9.72, 2.38, GREEN)
bullets(s, 7.1, 4.35, 5.45, 1.8, [
    "Task defines which information and gradient paths must survive.",
    "Matched baselines make every extra state earn its place.",
], 18)

# 4
s = prs.slides.add_slide(blank); title(s, "Review 1: is the scientific position credible?")
for i, (lab, sub) in enumerate([
    ("Purpose", "Test novelty and define a safe claim"),
    ("Method", "Protocol search + screening + full-text extraction + snowballing"),
    ("Evidence", "37 extracted studies across aggregate, mesoscopic, micro, and hybrid models"),
    ("Conclusion", "A narrow open intersection—not a universal priority claim"),
]):
    y = 1.25 + i * 1.18
    pill(s, .72, y, 1.65, lab, TEAL if i == 3 else PALE, TEAL)
    box(s, 2.6, y-.08, 9.95, .68, sub, LIGHT, PALE, 18, i == 3, PP_ALIGN.LEFT)
text(s, .8, 6.3, 11.8, .42, "The strongest threat is empirical: lane/storage aggregates may already provide equally useful control information.", 18, RED, True, PP_ALIGN.CENTER)

# 5
s = prs.slides.add_slide(blank); title(s, "The defensible gap and contribution")
for i, (head, body, color) in enumerate([
    ("Inherited", "LWR/LTM propagation · node constraints · continuous service · route commodities", BLUE),
    ("Adapted", "Divisible mass + persistent semantic order · localized lane/storage detail", AMBER),
    ("Potentially novel", "Selective-discreteness stack + matched evidence of better decision gradients", GREEN),
    ("Unsafe", "“First differentiable simulator” · universal mesoscopic fidelity · usefulness from finite gradients", RED),
]):
    y = 1.15 + i * 1.28
    pill(s, .65, y, 2.2, head, RGBColor(243,246,248), color)
    box(s, 3.05, y-.08, 9.55, .72, body, LIGHT, color, 18, False, PP_ALIGN.LEFT)
text(s, .75, 6.42, 11.85, .38, "Claim: preserve selected local order/access information in a continuous-mass derivative path—and test whether it matters.", 18, NAVY, True, PP_ALIGN.CENTER)

# 6
s = prs.slides.add_slide(blank); title(s, "Baseline ladder: make the contribution falsifiable")
rows = [
    ("B0", "Homogeneous macro", "Seo differentiable LTM + van de Weg signal allocation → matched effective-service/TST macro; controls implementation-only gain."),
    ("B1", "Movement / partial FIFO", "Wei turn queues + Wright partial FIFO → adopt one movement-queue/restriction formulation; tests whether movement structure suffices."),
    ("B2", "Lane + finite storage", "Gong multilane allocation + Liu/Wong/Gao storage → transcribe a declared aggregate lane/bay model; tests storage sufficiency."),
    ("B3", "Selective semantic order", "Individual-tracking LTM + differentiable-simulation precedents → project-owned composition; tests exact realized accessibility."),
    ("R1 / R2", "Ordered / microscopic reference", "Individual/event mesoscopic models + SUMO → reproduce bounded order cases and directly reuse SUMO; tests external transfer."),
]
for i, (code, role, q) in enumerate(rows):
    y = 1.16 + i * 1.02
    pill(s, .55, y, 1.2, code, TEAL if code == "B3" else PALE, TEAL)
    box(s, 1.95, y-.05, 2.75, .88, role, LIGHT, PALE, 18, code == "B3", PP_ALIGN.LEFT)
    box(s, 4.92, y-.05, 7.85, .88, q, LIGHT, PALE, 18, False, PP_ALIGN.LEFT)
text(s, .75, 6.42, 11.85, .35, "Direct adoption is limited: equations anchor B0–B2, composition defines B3, and maintained software supports the external evaluator.", 18, NAVY, True, PP_ALIGN.CENTER)

# 7
s = prs.slides.add_slide(blank); title(s, "Why these three mesoscopic ingredients?")
items = [
    ("ORDER", "FIFO rank / semantic ledger", "Literature: individual LTM + LIFT retain realized order.\nEvolution: adapt identity/order into a continuous-mass semantic ledger."),
    ("ACCESS", "Lane intent ≠ realized reachability", "Literature: Wright + Wei establish restricted movement service.\nEvolution: modify and glue it to ledger-conditioned exact accessibility."),
    ("STORAGE", "Finite bay occupancy + spillback", "Literature: Liu–Chang, Wong–Lee, Gao + Gong establish storage and lane effects.\nEvolution: adopt these locally and couple them to aggregate LTM."),
]
for i, (head, mid, tail) in enumerate(items):
    x = .65 + i * 4.18
    box(s, x, 1.25, 3.72, .73, head, PALE, TEAL, 20, True)
    text(s, x+.12, 2.18, 3.48, .68, mid, 18, NAVY, True, PP_ALIGN.CENTER)
    text(s, x+.13, 3.0, 3.46, 2.15, tail, 18, INK, False, PP_ALIGN.CENTER)
    if i < 2: text(s, x+3.78, 2.65, .55, .5, "+", 27, TEAL, True, PP_ALIGN.CENTER)
box(s, .82, 5.55, 11.7, .82, "No ingredient is claimed as invented alone: the novelty candidate is their selective continuous-mass composition plus the matched derivative/usefulness test.", LIGHT, GREEN, 18, True)

# 8
s = prs.slides.add_slide(blank); title(s, "Scope boundary by example")
text(s, .7, 1.1, 5.9, .4, "IN SCOPE", 19, GREEN, True, PP_ALIGN.CENTER)
text(s, 6.75, 1.1, 5.9, .4, "OUT OF SCOPE", 19, RED, True, PP_ALIGN.CENTER)
for i, (yes, no) in enumerate([
    ("Continuous vehicle-equivalent mass", "Integerizing vehicles or microscopic car-following"),
    ("Lane-local order/access near a node", "Full-network lane expansion or lane-changing dynamics"),
    ("Fixed routes, permissions, phase order", "Dynamic route choice or phase-sequence decisions"),
    ("Effective green/service allocation", "Executable cycle, yellow/all-red, offsets, actuation"),
    ("Matched gradients + richer evaluation", "Claiming usefulness from differentiability alone"),
]):
    y = 1.63 + i * .92
    box(s, .68, y, 5.78, .62, "✓  " + yes, LIGHT, GREEN, 17, False, PP_ALIGN.LEFT)
    box(s, 6.86, y, 5.78, .62, "×  " + no, LIGHT, RED, 17, False, PP_ALIGN.LEFT)
text(s, .8, 6.42, 11.75, .33, "The project tests a selective representation—not a universally richer traffic simulator.", 18, NAVY, True, PP_ALIGN.CENTER)

# 9
s = prs.slides.add_slide(blank); title(s, "Review 2: which downstream task can expose the mechanism?")
box(s, .62, 1.2, 3.35, 1.0, "Primary family\nfixed-time green allocation", PALE, TEAL, 18, True)
box(s, 4.98, 1.2, 3.35, 1.0, "First meaning\ncontinuous effective service", PALE, GREEN, 18, True)
box(s, 9.34, 1.2, 3.35, 1.0, "Deferred meaning\nexecutable signal timing", LIGHT, AMBER, 18, True)
arrow(s, 3.98, 1.7, 4.95, 1.7); arrow(s, 8.34, 1.7, 9.31, 1.7)
bullets(s, .85, 2.82, 11.65, 2.9, [
    "Candidate families reviewed: signal allocation, gradient perturbation, corridor coupling, multi-fidelity optimization, and robust evaluation.",
    "First target: allocate service among fixed protected movement groups at one junction.",
    "Why first: low-dimensional action, direct capacity path, and clean matching across model representations.",
    "Current state: near the end of task definition, still open to implementation-driven calibration and approved refinements.",
], 18)

# 10
s = prs.slides.add_slide(blank); title(s, "Concrete task: symmetric four-leg service allocation")
cx, cy = 3.9, 3.65
for x, y, w, h in [(cx-1.0, .95, 2.0, 5.4), (.65, cy-.72, 6.5, 1.44)]:
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(226,231,234); sh.line.fill.background()
box(s, cx-.65, cy-.55, 1.3, 1.1, "Junction", WHITE, NAVY, 16, True)
text(s, .72, 1.0, 1.45, .35, "West", 15, MID); text(s, 5.65, 1.0, 1.45, .35, "East", 15, MID, False, PP_ALIGN.RIGHT)
text(s, 3.15, .68, 1.5, .32, "North", 15, MID, False, PP_ALIGN.CENTER); text(s, 3.15, 6.43, 1.5, .32, "South", 15, MID, False, PP_ALIGN.CENTER)
for i,(lab,col) in enumerate([("g₁  N–S through/right",BLUE),("g₂  N–S protected left",TEAL),("g₃  E–W through/right",AMBER),("g₄  E–W protected left",GREEN)]):
    box(s, 7.75, 1.15+i*1.0, 4.85, .66, lab, LIGHT, col, 17, True, PP_ALIGN.LEFT)
bullets(s, 7.75, 5.36, 4.85, 1.2, ["Fixed topology and routes", "Clearance is not a fifth group"], 16)

# 11
s = prs.slides.add_slide(blank); title(s, "Derivative contract: what do x, f, and df/dx mean?")
p = formula("task_contract", [
    r"x\equiv g=(g_1,g_2,g_3,g_4)^\mathsf{T},\quad g_p\geq g_p^{\min},\quad \sum_{p=1}^{4}g_p=G",
    r"f(x;\xi)\equiv J(g;\xi)=\Delta t\sum_{k=0}^{K-1}\left(N_k^{\rm src}+N_k^{\rm net}\right)",
    r"\frac{df}{dx}=\nabla_gJ,\qquad [J]={\rm veh}\!\cdot\!{\rm s}",
    r"D(g;\xi)=J(g;\xi)-J^{\rm ff}(\xi),\qquad \nabla_gD=\nabla_gJ",
], 12.2, 22)
picture(s, p, .55, 1.15, 12.2, 3.65)
box(s, .85, 5.1, 3.45, .9, "x\nphysical service fractions", PALE, TEAL, 17, True)
box(s, 4.93, 5.1, 3.45, .9, "f\nclearance-complete total system time", PALE, TEAL, 17, True)
box(s, 9.0, 5.1, 3.45, .9, "df/dx\nvehicle-seconds per service fraction", PALE, TEAL, 17, True)
text(s, .8, 6.35, 11.8, .38, "Nnet counts every vehicle-equivalent inside the network—moving or queued; it is not a queue-only objective.", 17, NAVY, True, PP_ALIGN.CENTER)

# 12
s = prs.slides.add_slide(blank); title(s, "V0–V4: add one causal challenge at a time")
variants = [
    ("V0", "exclusive lanes", "macro equivalence", "none"),
    ("V1", "shared lane", "head-of-line order", "+ ordered prefix"),
    ("V2", "finite left bay", "storage + spillback", "+ bay state"),
    ("V3", "same finite bay", "order-conditioned access", "+ exact access"),
    ("V4", "held-out family", "robustness by stratum", "+ scenario bank"),
]
for i,(v,layout,intent,added) in enumerate(variants):
    y=1.12+i*1.04
    pill(s,.55,y,1.0,v,TEAL if v=="V3" else PALE,TEAL)
    box(s,1.72,y-.05,2.55,.64,layout,LIGHT,PALE,17,v=="V3",PP_ALIGN.LEFT)
    box(s,4.48,y-.05,4.35,.64,intent,LIGHT,PALE,17,False,PP_ALIGN.LEFT)
    box(s,9.05,y-.05,3.72,.64,added,PALE if v!="V0" else LIGHT,TEAL,17,True,PP_ALIGN.LEFT)
text(s,.72,6.43,11.9,.35,"V3 is decisive only after V0 verifies equivalence and V2 isolates storage without binding order.",17,NAVY,True,PP_ALIGN.CENTER)

# 13
s = prs.slides.add_slide(blank); title(s, "Which literature-backed baseline enters each variant?")
headers=[("Variant",.55,1.0),("Required comparison",1.68,4.0),("Method / engineering affinity",5.9,6.88)]
for h,x,w in headers: box(s,x,1.1,w,.55,h,NAVY,NAVY,16,True)
data=[
    ("V0","B0–B3","Seo: differentiable LTM patterns · van de Weg/COS: signal-node equations"),
    ("V1","B0–B3 + R1","Wei: turn queues · Wright: partial FIFO · individual LTM: ordered reference"),
    ("V2","B0, B2, B3 + R1","Gong: multilane allocation · Liu/Wong/Gao: finite storage and blockage"),
    ("V3","B0–B3 + R1/R2","Same aggregate controls · selective order · SUMO external evaluation"),
    ("V4","B0–B3 + R2","Matched scenario bank · replicated SUMO seeds · optimizer kept separate"),
]
for i,(v,roles,aff) in enumerate(data):
    y=1.82+i*.91
    box(s,.55,y,1.0,.62,v,PALE,TEAL,16,True)
    box(s,1.68,y,4.0,.62,roles,LIGHT,PALE,16,False,PP_ALIGN.LEFT)
    box(s,5.9,y,6.88,.62,aff,LIGHT,PALE,15,False,PP_ALIGN.LEFT)
text(s,.72,6.48,11.9,.3,"Adoption unit: equations for B0–B2; documented composition for B3; direct software reuse for SUMO.",16,NAVY,True,PP_ALIGN.CENTER)

# 14
s = prs.slides.add_slide(blank); title(s, "General gradient path: control → physics → time")
nodes=[("g","service\nallocation"),("cₘ(g)","movement\ncapacity"),("yₘ,ₖ","realized\ndischarge"),("Xₖ₊₁","next traffic\nstate"),("Nsrc,Nnet","system\noccupancy"),("f","total system\ntime")]
for i,(sym,lab) in enumerate(nodes):
    x=.35+i*2.12
    box(s,x,1.45,1.72,1.02,sym,PALE,TEAL,20,True)
    text(s,x,2.58,1.72,.72,lab,16,MID,False,PP_ALIGN.CENTER)
    if i<len(nodes)-1: arrow(s,x+1.75,1.96,x+2.08,1.96,TEAL,2)
p=formula("sensitivity",[
    r"X_{k+1}=F_k(X_k,g;\xi),\qquad S_k=\frac{\partial X_k}{\partial g}",
    r"S_0=0,\qquad S_{k+1}=A_kS_k+B_k,\quad A_k=\frac{\partial F_k}{\partial X_k},\quad B_k=\frac{\partial F_k}{\partial g}",
    r"\nabla_g f=\sum_{k=0}^{K-1}S_k^\mathsf{T}\nabla_{X_k}\ell_k",
],11.7,21)
picture(s,p,.8,3.55,11.7,2.25)
text(s,.8,6.35,11.8,.35,"Aₖ transports earlier control effects; Bₖ adds the current step’s direct control effect. Reverse-mode AD computes the same chain.",17,NAVY,True,PP_ALIGN.CENTER)

# 15
s = prs.slides.add_slide(blank); title(s, "Derivative validity and usefulness are different tests")
p=formula("validation",[
    r"\mathrm{D}J(g;\xi)[d]=\nabla_gJ(g;\xi)^\mathsf{T}d,\qquad \mathbf{1}^\mathsf{T}d=0",
    r"\widehat{\mathrm{D}}_hJ[d]=\frac{J(g+hd;\xi)-J(g-hd;\xi)}{2h}",
    r"\Delta J_R^{(i)}=J_R(g+\delta d_i;\xi)-J_R(g;\xi)",
],11.3,23)
picture(s,p,.75,1.2,11.85,2.55)
box(s,.7,4.2,3.65,1.05,"1  Internal correctness\nAD ≈ central difference",PALE,BLUE,17,True)
box(s,4.83,4.2,3.65,1.05,"2  Direction quality\nsign + ranking",PALE,AMBER,17,True)
box(s,8.96,4.2,3.65,1.05,"3  External usefulness\nrealized improvement in R",PALE,GREEN,17,True)
text(s,.75,5.75,11.85,.68,"Differentiate continuous mass only within a stable active-set / prefix signature. Report event-boundary cases separately; do not invent derivatives through labels, permutations, or branch changes.",17,NAVY,True,PP_ALIGN.CENTER)

# 16
s = prs.slides.add_slide(blank); title(s, "Representative V1 path: shared-lane order changes future service")
p=formula("v1",[r"y_{m,k}=\min\!\left\{E_{m,k}(X_k;\xi_{\rm ord}),\ R_{m,k},\ c_m(g)\Delta t\right\}"],11.5,24)
picture(s,p,.7,1.08,11.9,.95)
labels=[("g","service"),("yₖ","discharge"),("residual head","continuous mass"),("Eₖ₊₁","eligible prefix"),("yₖ₊₁","future discharge"),("f","system time")]
for i,(a,b) in enumerate(labels):
    x=.32+i*2.13
    box(s,x,2.45,1.75,.88,a,PALE,TEAL,18,True)
    text(s,x,3.4,1.75,.45,b,14,MID,False,PP_ALIGN.CENTER)
    if i<5: arrow(s,x+1.77,2.89,x+2.10,2.89)
box(s,.7,4.35,5.7,1.22,"Fixed structural scenario\nroute labels · movement labels · ordering skeleton",LIGHT,AMBER,17,True)
box(s,6.93,4.35,5.7,1.22,"Differentiated quantities\nprefix mass · partial removal · residual mass",LIGHT,GREEN,17,True)
text(s,.75,6.2,11.85,.45,"Aggregate-matched order pairs share totals and g, but their fixed order scenarios can produce different f and ∂f/∂g.",17,NAVY,True,PP_ALIGN.CENTER)

# 17
s = prs.slides.add_slide(blank); title(s, "Representative V3 path: order-conditioned bay accessibility")
p=formula("v3",[
    r"b_k=\mathcal{A}\!\left(L_k^{\rm ord},I_k^{\rm lane},r_k^{\rm bay}\right)",
    r"g\rightarrow y_k\rightarrow(O_{k+1}^{\rm bay},I_{k+1}^{\rm lane},L_{k+1}^{\rm ord})\rightarrow\mathcal{A}_{k+1}\rightarrow b_{k+1}\rightarrow y_{k+1}\rightarrow f",
],12.0,20)
picture(s,p,.55,1.08,12.2,1.7)
box(s,.62,3.25,3.7,1.25,"BAY STORAGE\noccupancy → receiving room",LIGHT,BLUE,17,True)
box(s,4.82,3.25,3.7,1.25,"REALIZED ACCESS\nwhich intended users reach bay",LIGHT,AMBER,17,True)
box(s,9.02,3.25,3.7,1.25,"ORDERED LEDGER\ncontinuous masses, fixed labels",LIGHT,GREEN,17,True)
arrow(s,4.33,3.88,4.79,3.88); arrow(s,8.53,3.88,8.99,3.88)
bullets(s,.82,5.02,11.6,1.25,[
    "B2 can share occupancy and receiving room, but lacks the exact ledger-conditioned accessibility path.",
    "The exact Jacobian of A depends on the approved operator design and must be fixed during development.",
],17)

# 18
s = prs.slides.add_slide(blank); title(s, "Where the task review stands")
box(s,.65,1.2,3.75,1.05,"DEFINED\nscientific input, objective, scope",PALE,GREEN,18,True)
box(s,4.8,1.2,3.75,1.05,"STRUCTURED\nV0–V4 + matched baseline ladder",PALE,GREEN,18,True)
box(s,8.95,1.2,3.75,1.05,"OPEN TO EVIDENCE\ncalibration + exact operators",PALE,AMBER,18,True)
bullets(s,.8,2.85,11.7,2.95,[
    "Framework-level gradient chain is complete and acts as a model-design anchor.",
    "Exact lane-allocation, storage, ordered-access, and service derivatives follow their milestone equations.",
    "Numerical geometry, demand regimes, tolerances, seeds, and budgets are calibrated after interfaces exist—then frozen before confirmation.",
    "The review is near its endpoint, but implementation evidence may trigger the smallest explicit refinement rather than silent adaptation.",
],18)
text(s,.8,6.3,11.75,.4,"Target conclusion remains conditional: learn where selective semantic order improves decisions—and where simpler aggregates are enough.",18,NAVY,True,PP_ALIGN.CENTER)

prs.core_properties.title = "From Selective Discreteness to a Testable Control Task"
prs.core_properties.subject = "Weekly literature-positioning and downstream-task review"
prs.core_properties.author = "diff_minimal_meso project"
# Enforce the requested 18 pt floor for all editable slide text. Formula graphics
# retain their own larger mathematical typesetting.
for slide in prs.slides:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is None or run.font.size.pt < 18:
                    run.font.size = Pt(18)
prs.save(OUT)
print(OUT)
